"""Genera docs/informe_final.docx (sobre la plantilla oficial de la catedra) y docs/informe_final.pdf.

Todo numero del informe se lee de los archivos generados por las ejecuciones (results/, data/, figures/).
Las afirmaciones cualitativas que dependen de los resultados se verifican con `claim(...)`: si un resultado
cambiara, el informe no se genera hasta revisar el texto.

Uso:
    python scripts/build_report.py                         # DOCX + PDF (el PDF requiere Microsoft Word)
    python scripts/build_report.py --repo-url https://github.com/USUARIO/REPO   # completa la URL y regenera
    python scripts/build_report.py --no-pdf                # solo DOCX
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from report_lib import (ReportBuilder, add_page_number_footer, duplicate_last_row, ent, ensure_outline_levels,  # noqa: E402
                        num, pct, remove_body_from, render, replace_text_everywhere, set_cell_text)

from charizard_ga.config import load_experiments, load_quality_criterion  # noqa: E402
from charizard_ga.data import TYPE_ORDER, load_catalog, load_metadata  # noqa: E402
from charizard_ga.fitness import FITNESS_VERSION, components, contributions, weighted_fitness  # noqa: E402
from charizard_ga.paths import CONFIG_DIR, DOCS_DIR, FIGURES_DIR, RESULTS_DIR, ROOT  # noqa: E402

PLACEHOLDER_REPO = "[[HL:[COMPLETAR: URL del repositorio público de GitHub]]]"


def claim(condition: bool, message: str):
    if not condition:
        raise AssertionError("El texto del informe ya no coincide con los resultados: " + message)


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


# ---------------------------------------------------------------------------
# Carga de evidencias
# ---------------------------------------------------------------------------
class Ctx:
    def __init__(self):
        self.cat = load_catalog()
        self.meta = load_metadata()
        self.exps = load_experiments()
        self.quality = load_quality_criterion()
        self.report_cfg = read_json(CONFIG_DIR / "report.json")
        self.ds_cfg = read_json(CONFIG_DIR / "data_sources.json")
        self.runs = read_csv(RESULTS_DIR / "summary_runs.csv")
        self.configs = {r["config_id"]: r for r in read_csv(RESULTS_DIR / "summary_configs.csv")}
        self.qa = {r["config_id"]: r for r in read_csv(RESULTS_DIR / "quality_assessment.csv")}
        self.gaps = read_csv(RESULTS_DIR / "ga_vs_optimum.csv")
        self.cross = {r["solucion_de"]: r for r in read_csv(RESULTS_DIR / "cross_profile_matrix.csv")}
        base = RESULTS_DIR / "baselines"
        self.opt = {r["perfil"]: r for r in read_csv(base / "exhaustive_optimum.csv")}
        self.top = read_csv(base / "exhaustive_top10.csv")
        self.cons = {r["perfil"]: r for r in read_csv(base / "conservative_optimum.csv")}
        self.methods = read_csv(base / "method_comparison.csv")
        self.greedy = read_csv(base / "greedy.csv")
        self.rs = read_csv(base / "random_search.csv")
        self.ex_info = read_json(base / "exhaustive_info.json")
        self.env = read_json(RESULTS_DIR / "environment.json")
        self.params = read_json(RESULTS_DIR / "parameters.json")
        self.legal = read_json(RESULTS_DIR / "legality_showdown.json") if (RESULTS_DIR / "legality_showdown.json").exists() else None
        self.legal_rows = read_csv(RESULTS_DIR / "legality_showdown.csv") if (RESULTS_DIR / "legality_showdown.csv").exists() else []
        self.validation = read_json(RESULTS_DIR / "validation_summary.json") if (RESULTS_DIR / "validation_summary.json").exists() else None
        self.main = {cid: read_json(RESULTS_DIR / "runs" / cid / "seed_42_summary.json") for cid in ("R1", "R2", "R3", "R4", "R5")}
        self.main_logs = {cid: read_csv(RESULTS_DIR / "runs" / cid / "seed_42_log.csv") for cid in self.main}
        self.learnset = read_csv(ROOT / "data" / "learnset_sv_charizard.csv")
        self.catalog_rows = read_csv(ROOT / "data" / "catalog_moves.csv")

    def tes(self, t: str) -> str:
        return self.cat.type_es[t]

    def types_es(self, types) -> str:
        return ", ".join(self.tes(t) for t in types)

    def mv_name(self, slug: str, en=False) -> str:
        m = self.cat.moves[self.cat.by_slug(slug)]
        return f"{m.nombre_es} ({m.nombre_en})" if en else m.nombre_es

    def moveset_es(self, slugs_text: str) -> str:
        return ", ".join(self.mv_name(s) for s in slugs_text.split("|"))


# ---------------------------------------------------------------------------
# Contenido
# ---------------------------------------------------------------------------
def build_content(c: Ctx, repo_url: str | None) -> ReportBuilder:
    R = ReportBuilder()
    cat, meta = c.cat, c.meta
    M = cat.M
    CM4 = math.comb(M, 4)
    counts = meta["conteos"]
    excl = counts["excluidos_por_motivo"]
    prof = c.exps["profiles"]
    pinfo = c.exps["profile_info"]
    ga = c.exps["ga_base"]
    n_runs = len(c.runs)
    hits_total = sum(1 for g in c.gaps if g["alcanza_optimo"] == "True")
    ex_secs = float(c.ex_info["segundos_enumeracion_y_componentes"])
    mean_eff = statistics.fmean(int(r["evaluaciones_efectivas"]) for r in c.runs if r["grupo"] == "principal")
    mean_req = statistics.fmean(int(r["evaluaciones_solicitadas"]) for r in c.runs if r["grupo"] == "principal")
    mean_time = statistics.fmean(float(r["tiempo_segundos"]) for r in c.runs if r["grupo"] == "principal")
    pk_commit = c.ds_cfg["pokeapi"]["commit"]
    sd_commit = c.ds_cfg["showdown"]["commit"]
    fecha_datos = meta["fecha_obtencion_utc"]["primera"][:10]
    repo_txt = repo_url if repo_url else PLACEHOLDER_REPO
    r1 = c.main["R1"]["mejor_solucion"]
    r1_slugs = "|".join(m["slug"] for m in r1["movimientos"])
    opt_r1 = float(c.opt["R1"]["optimo"])
    max_cov = int(c.ex_info["cobertura_maxima_tipos"])
    n_max_cov = int(c.ex_info["combinaciones_con_cobertura_maxima"])

    claim(hits_total == n_runs, "todas las corridas alcanzan el optimo")
    claim(sum(int(r["invalidos_registrados"]) for r in c.runs) == 0, "cero invalidos registrados")
    main_runs = [r for r in c.runs if r["grupo"] == "principal"]
    gen_lo = min(int(r["generacion_encontrada"]) for r in main_runs)
    gen_hi = max(int(r["generacion_encontrada"]) for r in main_runs)
    gmeans = [float(c.configs[k]["generacion_encontrada_media"]) for k in ("R1", "R2", "R3", "R4", "R5")]
    div = {k: float(c.configs[k]["distintas_poblacion_final_media"]) for k in c.configs}
    claim(all(8 <= div[k] <= 12 for k in ("R1", "R2", "R3", "R4", "R5", "OP-1P")), "unas 10 combinaciones distintas con torneo")
    claim(4 <= div["OP-RUL"] / div["R1"] <= 7, "ruleta ~5 veces mas diversidad")
    claim(1.7 <= div["OP-MUT20"] / div["R1"] <= 2.7, "mutacion 0,20 ~duplica diversidad")
    rul_log = read_csv(RESULTS_DIR / "runs" / "OP-RUL" / "seed_42_log.csv")
    rul_ratio = max(float(r["aptitud_max"]) / float(r["aptitud_min"]) for r in rul_log)
    claim(all("focus-punch" in r["moveset_slugs"] for r in c.runs), "Puno Certero en todas las soluciones")
    claim(M == 53 and CM4 == 292_825, "tamano del catalogo")

    # --- Resumen y palabras clave se escriben en la plantilla (ver fill_template) ---------------

    R.page_break()
    R.toc()
    R.page_break()

    # =========================================================================================
    R.h1("1. Introducción")
    R.h2("1.1. Problema")
    R.p("El trabajo aborda la **Propuesta 1** aprobada en la entrega parcial: encontrar, para el Pokémon **Charizard** "
        "(tipos Fuego y Volador), una combinación de **cuatro movimientos ofensivos diferentes y permitidos** que maximice "
        "un equilibrio entre cobertura de tipos, potencia, precisión, STAB (bonificación por coincidencia de tipo) y diversidad. "
        "Charizard se mantiene fijo durante todo el problema: lo único que se optimiza es su *moveset*.")
    R.p("**Contexto competitivo.** Se toma como referencia un formato individual en el que cada jugador prepara un equipo de seis "
        "Pokémon y, tras ver el equipo rival, elige tres para el combate. La configuración registrada de cada Pokémon —incluidos "
        "sus cuatro movimientos— queda fija durante una serie de combates. Por eso conviene que Charizard tenga una respuesta "
        "ofensiva razonable frente a rivales variados, en lugar de un moveset óptimo contra un único rival. El formato sirve "
        "solo de contexto: no se optimiza el equipo de seis ni se cambian movimientos frente a cada rival.")
    R.h2("1.2. Ejemplo")
    fire = [cat.by_slug(s) for s in ("flamethrower", "fire-blast", "heat-wave", "overheat")]
    part = [cat.by_slug(s) for s in ("flamethrower", "air-slash", "dragon-pulse", "solar-beam")]
    cf, cp = components(fire, cat), components(part, cat)
    ff, fp = weighted_fitness(cf, prof["R1"]), weighted_fitness(cp, prof["R1"])
    claim(len(cf.covered) == 4 and len(cp.covered) == 9 and fp > ff, "ejemplo de introduccion")
    R.p(f"Un moveset formado solo por ataques de Fuego —Lanzallamas, Llamarada, Onda Ígnea y Sofoco— aprovecha el STAB en los "
        f"cuatro movimientos y tiene buena potencia, pero solo es superefectivo contra {len(cf.covered)} de los 18 tipos "
        f"({c.types_es(cf.covered)}). El ejemplo de la entrega parcial —Lanzallamas, Tajo Aéreo, Pulso Dragón y Rayo Solar— "
        f"resigna STAB y algo de precisión, pero cubre {len(cp.covered)} tipos. Con la función de aptitud y el perfil de pesos "
        f"base que se definen en la Sección 2, el primero obtiene F = {num(ff, 2)} y el segundo F = {num(fp, 2)}: los criterios "
        "entran en conflicto y el mejor equilibrio no surge de elegir individualmente los movimientos «mejores».")
    R.h2("1.3. Relevancia y características")
    R.p("Desde el punto de vista del jugador, el problema representa la preparación previa de un Pokémon que debe ser útil "
        "frente a rivales diversos. Desde el punto de vista de la asignatura, es un problema de **optimización combinatoria "
        "con objetivos en conflicto** —cobertura frente a potencia, potencia frente a precisión, STAB frente a diversidad— "
        "análogo a la selección de características o a la configuración de componentes, donde el valor de un elemento depende "
        "de los demás (dos movimientos del mismo tipo aportan potencia pero no cobertura nueva). Las restricciones son duras: "
        "cuatro movimientos distintos y todos aprendibles por Charizard en la versión definida.")
    R.h2("1.4. Diferencia con un problema ya resuelto")
    R.p("El anexo de la consigna incluye «Seleccionar el pokemon más apto para ganar la lucha contra otro pokemon determinado». "
        "Este trabajo es distinto en tres aspectos: (a) el Pokémon no se elige, **Charizard es fijo**; (b) lo que se optimiza "
        "son sus **cuatro movimientos**, no la identidad del Pokémon; y (c) el objetivo es una **cobertura general** frente a los "
        "18 tipos individuales, no vencer a un rival concreto.")
    R.h2("1.5. Alcance")
    R.p("Se respetan las decisiones de la entrega parcial: Charizard normal (Fuego/Volador); individuos de exactamente cuatro "
        "movimientos; catálogo de una única versión (Pokémon Escarlata/Púrpura); cobertura medida contra los 18 tipos "
        "individuales y no contra combinaciones de doble tipo ni Pokémon concretos; solo movimientos de daño directo; y "
        "aptitud 0 para individuos inválidos. Quedan fuera la simulación completa de combate, objetos, EV, naturalezas, "
        "habilidades, clima, teracristalización y cambios de Pokémon.")
    R.h2("1.6. Justificación del Algoritmo Genético y tamaño del espacio de búsqueda")
    R.p(f"Con el catálogo definitivo de **M = {M}** movimientos, una solución es un subconjunto de 4 elementos, por lo que el "
        f"espacio tiene **C({M},4) = {ent(CM4)}** movesets. La propia entrega parcial reconoció que una búsqueda exhaustiva es "
        f"posible para catálogos pequeños, y efectivamente lo es: enumerar y evaluar todas las combinaciones con la misma "
        f"función tomó {num(ex_secs, 2)} s en el equipo de prueba. No se afirma, por lo tanto, que el problema sea imposible de "
        "resolver con software tradicional, ni se agrandó artificialmente el catálogo.")
    R.p("El AG se justifica como **método de búsqueda y optimización combinatoria guiado por una función de aptitud con "
        "objetivos en conflicto**: el cromosoma se codifica de forma directa (cuatro genes), la calidad se mide sin datos "
        "etiquetados —por eso no corresponde una red neuronal— y la cruza permite recombinar buenos subconjuntos de movimientos. "
        "La enumeración exhaustiva se usa como **referencia exacta** para medir la calidad del AG, no como reemplazo de su "
        "implementación. La Discusión analiza con honestidad qué aporta el AG en un espacio de este tamaño.")

    # =========================================================================================
    R.h1("2. Elementos del Trabajo y Metodología")
    R.h2("2.1. Fuentes y versión de los datos")
    R.p(f"Se eligió el grupo de versiones **Pokémon Escarlata/Púrpura** (`scarlet-violet`, generación IX), compatible con el "
        "formato individual descripto. Las fuentes son **comunitarias** (no oficiales de The Pokémon Company ni de Game Freak) "
        "y se fijaron por commit para que la preparación sea reproducible:")
    R.bullets([
        f"**PokéAPI** (fuente principal): volcado estático `api-data` que respalda la API REST pública https://pokeapi.co, "
        f"commit `{pk_commit[:12]}`. Se usan el learnset de Charizard (pokemon/6), los {counts['learnset_total']} movimientos "
        f"aprendibles en Escarlata/Púrpura y los 18 tipos.",
        f"**Pokémon Showdown** (verificación cruzada): `data/moves.ts`, `data/learnsets.ts` y `data/typechart.ts` del commit "
        f"`{sd_commit[:12]}`. Se usa para contrastar learnset, potencia, precisión, tipo, categoría y tabla de tipos, y para "
        "extraer banderas mecánicas descriptivas (carga, recarga, retroceso, etc.).",
        f"Fecha de obtención: {fecha_datos} (UTC). Tiempo de adquisición y preparación con descarga completa: "
        f"{num(meta['tiempo_ultima_descarga']['segundos_totales'], 1)} s para "
        f"{meta['tiempo_ultima_descarga']['archivos_descargados']} archivos; este tiempo es independiente del tiempo de "
        "búsqueda (una vez preparado el dataset, el proyecto no vuelve a descargar nada).",
    ])
    R.p("Los atributos de cada movimiento se reconstruyen **para Escarlata/Púrpura** a partir del historial `past_values` de "
        "PokéAPI (cada entrada guarda el valor previo a un cambio en un grupo de versiones posterior), de modo que no se "
        "mezcla la legalidad de una versión con atributos de otra. El dataset definitivo, sus metadatos y un registro de "
        "fuentes (`data/FUENTES_DATOS.md`) se guardan en el repositorio, por lo que el proyecto funciona sin conexión.")
    R.h2("2.2. Preparación del catálogo")
    learn_total = counts["learnset_total"]
    R.p(f"El learnset de Charizard en Escarlata/Púrpura contiene **{learn_total}** movimientos aprendibles por nivel, MT o "
        f"huevo (no hay tutores en esta versión para Charizard según la fuente). Se aplicaron reglas de exclusión documentadas, "
        f"en orden y sin selección manual de movimientos (@tab:exclusiones). El resultado es el **catálogo definitivo de M = {M}** "
        "movimientos (Anexo B).")
    excl_examples = {}
    for r in c.learnset:
        if r["motivo_exclusion"]:
            excl_examples.setdefault(r["motivo_exclusion"], []).append(r["nombre_es"])
    rule_txt = {
        "ESTADO": "Movimiento de estado (no produce daño directo).",
        "POTENCIA_NO_NUMERICA": "Potencia no numérica: daño fijo, variable o dependiente del rival u objeto.",
        "PRECISION_NO_NUMERICA": "Precisión especial no numérica («nunca falla»); no se convierte a 100.",
        "GOLPES_MULTIPLES": "Golpes múltiples (la potencia por golpe no representa el daño total).",
        "POTENCIA_CERO": "Potencia 0 informada por la fuente.",
        "OHKO": "Movimiento de KO en un golpe.",
    }
    rows = []
    for code in ("ESTADO", "POTENCIA_NO_NUMERICA", "PRECISION_NO_NUMERICA", "GOLPES_MULTIPLES"):
        rows.append([rule_txt[code], str(excl.get(code, 0)), ", ".join(sorted(excl_examples.get(code, []))[:6]) +
                     (" …" if len(excl_examples.get(code, [])) > 6 else "") or "—"])
    rows.append(["**Total excluido / catálogo definitivo**", f"**{learn_total - M} / {M}**", f"C({M},4) = {ent(CM4)}"])
    R.table("exclusiones", f"Reglas de exclusión aplicadas al learnset de Escarlata/Púrpura ({learn_total} movimientos).",
            ["Regla (se aplica en orden)", "Cant.", "Movimientos excluidos (ejemplos)"], rows, [4.6, 1.0, 6.6], ["l", "c", "l"])
    R.p("**Movimientos con efectos que el modelo ignora.** Los movimientos de daño directo con carga (Vuelo, Excavar, Rayo Solar), "
        "recarga (Hiperrayo, Gigaimpacto, Anillo Ígneo), fallo condicional (Puño Certero, que falla si el usuario recibe daño "
        "antes de actuar), retroceso (Envite Ígneo, Doble Filo, Derribo), autodebilitamiento (Sofoco), bloqueo (Enfado), potencia "
        "condicional o tipo variable (Acróbata, Fachada, Meteorobola, Teraexplosión, Cólera Ardiente) **permanecen en el "
        "catálogo**, porque producen daño directo y tienen potencia y precisión numéricas. Sus efectos se documentan en columnas "
        "descriptivas del catálogo, pero la función aprobada usa la **potencia base sin transformarla en daño simulado**. "
        "Consecuencia: el modelo sobrevalora los movimientos cuyo costo real (un turno extra o un posible fallo) no se mide. "
        f"Para cuantificar este efecto se definió, con una regla previa a los experimentos, un **subconjunto conservador** de "
        f"{counts['subconjunto_conservador_M']} movimientos que excluye carga, recarga y fallo condicional; se usa solo como "
        "análisis de sensibilidad (Sección 3.6).")
    v = meta["verificacion_showdown"]
    only_sd = v["learnset_solo_en_showdown_detalle"]
    claim(not v["tabla_tipos_diferencias"] and not v["atributos_distintos_en_catalogo"] and not v["learnset_solo_en_pokeapi"],
          "verificacion cruzada sin diferencias")
    claim(list(only_sd) == ["falseswipe"], "unica diferencia de learnset")
    R.p(f"**Validación y verificación cruzada.** El script de preparación verifica IDs únicos, tipos reconocidos, potencias y "
        f"precisiones numéricas en rango, ausencia de duplicados y al menos cuatro movimientos válidos. Contra Showdown: la tabla "
        f"de tipos 18×18 coincide en las 324 celdas; los {M} movimientos del catálogo coinciden en potencia, precisión, tipo y "
        f"categoría; y todos los movimientos de PokéAPI figuran en el learnset de Gen 9 de Showdown. La única diferencia es "
        f"Falso Tortazo, que Showdown registra solo por MT de Charmander y Charmeleon (herencia por evolución); como el alcance "
        "declarado es el learnset propio de Charizard en PokéAPI, se informa pero no se incorpora.")
    if c.legal:
        claim(c.legal["movesets_validos"] == c.legal["movesets_validados"], "todos los movesets resultantes validos")
        R.p(f"**Coexistencia de los cuatro movimientos.** El catálogo garantiza la pertenencia *individual* de cada movimiento. "
            f"Para verificar que los cuatro puedan coexistir en un mismo Charizard, todos los movesets resultantes "
            f"({c.legal['movesets_validados']} combinaciones distintas: mejores de cada corrida, óptimos, heurísticas y ejemplo) se "
            f"validaron con el validador de equipos de Pokémon Showdown en el formato «[Gen 9] Anything Goes» (regla *Obtainable*), "
            f"commit `{sd_commit[:12]}`: {c.legal['movesets_validos']} de {c.legal['movesets_validados']} son válidos, y "
            f"{c.legal['controles_rechazados']} controles negativos (un movimiento no aprendible y uno repetido) fueron rechazados "
            "(Sección 3.7). Para cualquier otro moveset del catálogo la coexistencia no se verificó individualmente.")
    R.h2("2.3. Tabla de efectividad de tipos")
    R.p("Se guarda la tabla completa de los 18 tipos con **orientación explícita: filas = tipo atacante, columnas = tipo "
        "defensor** (`data/type_chart.csv`, @fig:tipos). Los multiplicadores válidos son 0, ½, 1 y 2. Las pruebas verifican "
        "casos asimétricos conocidos (Tierra→Fuego = 2 pero Fuego→Tierra = 1; Eléctrico→Tierra = 0 pero Tierra→Eléctrico = 2) y "
        "las ocho inmunidades de la tabla moderna.")
    R.figure("tipos", FIGURES_DIR / "fig_tabla_tipos.png",
             "Tabla de efectividad usada por el modelo (Gen 9). Filas: tipo del movimiento atacante; columnas: tipo defensor. "
             "La cobertura C(x) cuenta las columnas para las que algún movimiento del moveset tiene una celda azul (×2).", 11.0)

    R.h2("2.4. Cromosoma y restricciones")
    R.p("Cada individuo es una **lista de cuatro identificadores de movimientos distintos** del catálogo, por ejemplo "
        "[53, 403, 406, 76] = [Lanzallamas, Tajo Aéreo, Pulso Dragón, Rayo Solar]. El orden no cambia el significado ni la "
        "aptitud; por eso se usa una **clave ordenada** del moveset para detectar combinaciones equivalentes, medir la "
        "diversidad de la población y guardar en caché la aptitud, sin alterar el orden de los genes que manipulan los "
        "operadores. Restricciones duras: longitud exactamente 4, sin IDs repetidos y todos pertenecientes al catálogo. "
        "Se permiten movimientos diferentes del mismo tipo. Un individuo que no cumple recibe **aptitud 0** (penalización dura, "
        "como en el ejemplo de la teoría, pág. 20).")
    R.h2("2.5. Función de aptitud")
    R.p("Se implementó exactamente la familia de funciones aprobada, a maximizar:")
    R.eq("F(x) = 100 × ( w_{C}·C(x) + w_{P}·P(x) + w_{A}·A(x) + w_{S}·S(x) + w_{D}·D(x) )")
    R.p("con los cinco componentes en [0, 1]:")
    R.bullets([
        "**C(x)**, cobertura: cantidad de tipos defensores individuales para los que existe al menos un movimiento del moveset "
        "con multiplicador mayor que 1, dividida por 18 (cada tipo cuenta una sola vez).",
        f"**P(x)**, potencia: potencia base promedio de los cuatro movimientos dividida por la **potencia máxima del catálogo "
        f"definitivo ({cat.power_max})**. La normalización es fija para todas las generaciones, corridas y perfiles; no se "
        "recalcula con la población.",
        "**A(x)**, precisión: precisión promedio de los cuatro movimientos dividida por 100.",
        "**S(x)**, STAB: cantidad de movimientos de tipo Fuego o Volador dividida por 4. Representa la presencia de ataques "
        "que recibirían STAB; la función **no** aplica además un multiplicador de daño a P (definiciones separadas).",
        "**D(x)**, diversidad: cantidad de tipos distintos entre los cuatro movimientos dividida por 4.",
    ])
    R.p("En el código, el cálculo de componentes (`fitness.components`) está separado de la ponderación "
        "(`fitness.weighted_fitness`). La caché de aptitud pertenece a un único par (dataset, pesos): cada corrida crea su propio "
        f"evaluador, por lo que nunca se reutiliza una aptitud calculada con otros pesos u otro dataset. Versión de la función "
        f"registrada en cada log: «{FITNESS_VERSION}».")
    R.p("**Pesos.** Deben ser no negativos y sumar 1. La entrega parcial dejó sus valores para la implementación; se adoptó como "
        "perfil base w_{C} = 0,40, w_{P} = 0,20, w_{A} = 0,15, w_{S} = 0,15 y w_{D} = 0,10. Es una **decisión de modelado que prioriza la "
        "cobertura general** —el objetivo central de la propuesta—, no un valor oficial ni demostrado óptimo. Su sensibilidad se "
        "estudia con cinco perfiles (@tab:perfiles), que son entradas del experimento y representan distintas preferencias de "
        "un jugador, no verdades competitivas.")
    rows = []
    for pid in ("R1", "R2", "R3", "R4", "R5"):
        w = prof[pid]
        rows.append([pid, pinfo[pid]["nombre"], num(w.wC), num(w.wP), num(w.wA), num(w.wS), num(w.wD), pinfo[pid]["justificacion"]])
    R.table("perfiles", "Perfiles de pesos de las cinco configuraciones principales (todas suman 1).",
            ["Conf.", "Perfil", "w_{C}", "w_{P}", "w_{A}", "w_{S}", "w_{D}", "Justificación"], rows,
            [0.9, 1.6, 0.8, 0.8, 0.8, 0.8, 0.8, 5.6], ["c", "l", "c", "c", "c", "c", "c", "l"], font_pt=7.5)

    R.h2("2.6. Ejemplo numérico desarrollado")
    cp_contrib = contributions(cp, prof["R1"])
    R.p("Se evalúa con el perfil base el moveset del ejemplo de la entrega parcial. Todos los valores se obtuvieron con el mismo "
        "código que usan los experimentos (y se verifican contra un cálculo manual en `tests/test_fitness.py`).")
    rows = []
    for mid in part:
        m = cat.moves[mid]
        cov = sorted(cat.se_types[m.tipo], key=TYPE_ORDER.index)
        rows.append([f"{m.nombre_es} ({m.nombre_en})", m.tipo_es, str(m.potencia), str(m.precision),
                     "Sí" if m.stab else "No", c.types_es(cov) or "—"])
    R.table("ejemplo_movs", "Movimientos del ejemplo y tipos frente a los que cada uno es superefectivo.",
            ["Movimiento", "Tipo", "Pot.", "Prec.", "STAB", "Superefectivo contra"], rows, [3.2, 1.4, 0.9, 0.9, 0.9, 4.9],
            ["l", "l", "c", "c", "c", "l"])
    R.p(f"Tipos cubiertos (unión, cada uno una vez): {c.types_es(cp.covered)} → **C = {len(cp.covered)}/18 = {num(cp.C, 4)}**. "
        f"Potencia promedio (90+75+85+120)/4 = 92,5 → **P = 92,5/{cat.power_max} = {num(cp.P, 4)}**. Precisión promedio "
        f"(100+95+100+100)/4 = 98,75 → **A = {num(cp.A, 4)}**. Movimientos Fuego o Volador: 2 → **S = 2/4 = {num(cp.S, 2)}**. "
        f"Tipos distintos: 4 → **D = 4/4 = {num(cp.D, 2)}**.")
    R.eq(f"F = 100 × (0,40·{num(cp.C, 4)} + 0,20·{num(cp.P, 4)} + 0,15·{num(cp.A, 4)} + 0,15·{num(cp.S, 2)} + 0,10·{num(cp.D, 2)})"
         f" = {num(cp_contrib['C'], 2)} + {num(cp_contrib['P'], 2)} + {num(cp_contrib['A'], 2)} + {num(cp_contrib['S'], 2)} + "
         f"{num(cp_contrib['D'], 2)} = **{num(fp, 4)}**")

    R.h2("2.7. Algoritmo genético y operadores")
    R.p("Se usó **Python 3 con DEAP** (como en los ejemplos de la cátedra) y un **ciclo evolutivo explícito** "
        "(`src/charizard_ga/ga.py`) comentado paso a paso. Una corrida sigue la estructura vista en clase (teoría, pág. 12):")
    R.bullets([
        "**Población inicial** de N = 100 individuos, cada uno con cuatro movimientos muestreados **sin reemplazo** del catálogo "
        "(siempre válidos).",
        "**Evaluación** con caché por combinación; DEAP solo reevalúa los individuos cuya aptitud fue invalidada por la cruza o "
        "la mutación.",
        f"**Elitismo**: los {ga.elitism} mejores se **clonan** y pasan sin cambios a la generación siguiente (protegidos de la variación). "
        "`tools.selBest` se usa solo para esto.",
        f"**Selección** de N − {ga.elitism} padres por **torneo** de tamaño {ga.tournament_size} (`tools.selTournament`) o por **ruleta** "
        "(`tools.selRoulette`). Si la suma de aptitudes fuera 0, la ruleta se reemplaza por selección uniforme (`tools.selRandom`).",
        f"**Cruza por pareja** (padres consecutivos) con probabilidad {num(ga.cx_prob)}: **uniforme / binomial aleatoria** (cada "
        f"posición se intercambia con probabilidad {num(ga.cx_indpb, 1)}, `tools.cxUniform`) o **de un punto** (`tools.cxOnePoint`).",
        "**Reparación de duplicados** tras la cruza: si un hijo repite un movimiento, se conservan los genes no repetidos y cada "
        "posición repetida se completa con un movimiento ausente elegido al azar —con el generador de la corrida, por lo tanto "
        "reproducible— entre los genes de ambos padres; solo si no hubiera candidatos se usaría el resto del catálogo. El "
        "validador sigue asignando 0 a cualquier individuo que resultara inválido (en todas las corridas se registraron 0 "
        "inválidos).",
        f"**Mutación por individuo** con probabilidad {num(ga.mutation_prob)}: se reemplaza **un gen elegido al azar** por un "
        "movimiento permitido ausente en el individuo (distinto también del gen reemplazado).",
        "**Reemplazo generacional** con tamaño constante: nueva población = élites + hijos.",
        "**Mejor de la corrida**: se conserva una **copia** del mejor individuo de toda la corrida (solo se reemplaza ante una "
        "mejora estricta), diferenciada del mejor de cada generación.",
        f"**Paro**: al completar {ga.generations} generaciones evolutivas, además de la generación inicial 0 (único criterio).",
    ])
    rows = [
        ["Tamaño de población", str(ga.population_size), "—"],
        ["Generaciones", f"{ga.generations} (+ generación 0)", "criterio de paro"],
        ["Selección", f"torneo, k = {ga.tournament_size}", "por cada padre seleccionado"],
        ["Cruza", f"uniforme, p = {num(ga.cx_prob)}", "por pareja de padres"],
        ["Intercambio en cruza uniforme", num(ga.cx_indpb, 1), "por gen (posición) de una pareja cruzada"],
        ["Mutación", f"p = {num(ga.mutation_prob)}; 1 gen", "por individuo hijo"],
        ["Elitismo", f"{ga.elitism} individuos", "por generación"],
        ["Semillas", "42 a 51 (10 réplicas)", "una por corrida; 42 = corrida principal"],
    ]
    R.table("params", "Parámetros del AG comunes a las cinco configuraciones principales y alcance de cada probabilidad.",
            ["Parámetro", "Valor", "Se aplica"], rows, [4.0, 3.6, 4.6], ["l", "l", "l"])
    R.p("**Reproducibilidad.** Toda la aleatoriedad del AG (la interna de DEAP y la de los operadores propios) usa el módulo "
        "`random` de Python, que se inicializa con la semilla al comienzo de cada corrida (también se fija `numpy`, que no "
        "participa del AG). La búsqueda aleatoria usa un generador propio `random.Random(semilla)`. Las pruebas comprueban que la "
        "misma semilla y configuración producen exactamente el mismo log, excluyendo tiempos.")
    R.p("**Adaptaciones respecto de los ejemplos docentes.** Se revisaron los notebooks del ZIP antes de reutilizarlos: (1) en "
        "ellos «Ranking» se registra como `tools.selBest`, que elige determinísticamente a los mejores y no es una selección "
        "probabilística por ranking; aquí no se usa como selección. (2) El notebook del viajero usa una permutación con "
        "`cxOrdered` y `mutShuffleIndexes`; esos operadores no se copiaron porque intercambiar posiciones de un moveset no "
        "introduce una combinación nueva. (3) En los ejemplos la mutación se aplica a un solo individuo por generación; aquí la "
        "probabilidad es por individuo. (4) `creator.create` se invoca una sola vez (`ensure_deap_types`), para que las celdas "
        "puedan reejecutarse sin redefinir clases. (5) Cada corrida usa un `Toolbox` nuevo y clona padres, élites y mejor "
        "individuo, invalidando la aptitud de todo hijo modificado.")

    R.h2("2.8. Diseño experimental")
    R.p("La consigna exige al menos cinco corridas con algún parámetro distinto. Se ejecutaron **cinco configuraciones "
        "principales (R1–R5)** que varían solo el vector de pesos y mantienen constantes los parámetros del AG, más **tres "
        "configuraciones de operadores** con pesos base que cambian una sola decisión respecto de R1 (@tab:configs). Cada "
        f"configuración se ejecutó con **diez réplicas** (semillas 42 a 51, las mismas en todas). Las **corridas principales "
        "del informe son las de semilla 42**, elegida antes de observar resultados; las demás réplicas se usan solo para "
        f"analizar variabilidad. En total se realizaron **{n_runs} corridas**, todas completas (ninguna quedó sin ejecutar). "
        "No fue necesario ajustar los parámetros iniciales: una prueba previa no mostró problemas de ejecución.")
    rows = []
    for e in c.exps["configs"]:
        ov = ", ".join(f"{k} = {v}" for k, v in e.overrides.items()) or "ninguno (pesos del perfil)"
        rows.append([e.config_id, "Principal" if e.grupo == "principal" else "Operadores", e.profile, ov.replace("selection = roulette", "selección = ruleta")
                     .replace("crossover = one_point", "cruza = un punto").replace("mutation_prob = 0.2", "mutación = 0,20"), "42–51"])
    R.table("configs", "Configuraciones ejecutadas. Todas usan los parámetros de @tab:params salvo el cambio indicado.",
            ["Conf.", "Grupo", "Pesos", "Cambio respecto de la base", "Semillas"], rows, [1.6, 1.8, 1.1, 5.9, 1.6],
            ["l", "l", "c", "l", "c"])
    q = c.quality
    R.p("**Criterio de calidad** (definido antes de analizar resultados, `config/quality_criterion.json`). Se distinguen tres "
        "nociones: *solución válida* (cuatro movimientos distintos del catálogo); *mejor solución encontrada* (el mejor "
        "individuo de la corrida, sin implicar optimalidad); y *óptimo demostrado* (el máximo de F sobre las C(M,4) "
        "combinaciones, obtenido por enumeración exhaustiva con la misma función). Una corrida **alcanza el óptimo** si su mejor "
        f"solución es válida y su brecha absoluta es menor que {q['tolerancia_empate']:g}; es **aceptable** si la brecha relativa "
        f"es ≤ {pct(q['umbral_brecha_relativa'], 0)}. Una configuración es **satisfactoria** si al menos "
        f"{q['umbral_replicas_optimo']} de sus 10 réplicas alcanzan el óptimo y la brecha relativa media es ≤ "
        f"{pct(q['umbral_brecha_relativa'], 0)}. El criterio evalúa al AG como método de búsqueda bajo el modelo; no mide "
        "probabilidad de victoria.")
    R.p("**Registro por corrida.** Cada corrida guarda su configuración completa, semilla, versión del dataset "
        f"(`{cat.dataset_version}`) y versión de la aptitud; un **log CSV por generación** (incluida la 0) con mínimo, media, "
        "máximo, desvío, mejor acumulada, combinaciones distintas, evaluaciones solicitadas y efectivas, cruzas, mutaciones, "
        "reparaciones e inválidos; y un resumen JSON con el mejor individuo, la generación donde apareció, sus componentes, su "
        "aptitud propia y bajo el perfil base, los tipos cubiertos y no cubiertos, las generaciones ejecutadas, el motivo de paro "
        "y el tiempo medido. Los valores se guardan con precisión completa y se redondean solo para presentarlos.")
    R.p("**Comparación con otros métodos** (mismo catálogo y misma función): (a) **enumeración exhaustiva** de las "
        f"{ent(CM4)} combinaciones (óptimo exacto por perfil y brecha del AG); (b) **búsqueda aleatoria** de movesets válidos con "
        "las mismas semillas y dos presupuestos igualados al del AG en cada réplica: igual cantidad de evaluaciones *solicitadas* "
        "y igual cantidad de evaluaciones *efectivas* (combinaciones distintas); (c) **heurística voraz** reproducible: parte "
        "del moveset vacío y en cada paso agrega el movimiento con mayor aptitud parcial (casillas vacías aportan 0, mismas "
        "normalizaciones), desempatando por menor ID; y (d) **voraz + búsqueda local**: ascenso de colina que prueba todos los "
        "reemplazos de un movimiento y aplica la mejor mejora hasta no mejorar. Se distinguen candidatos generados, "
        "evaluaciones solicitadas, evaluaciones efectivas y reutilización por caché.")

    R.h2("2.9. Herramientas, lenguaje y librerías")
    env = c.env
    rows = [["Lenguaje", f"Python {env['python']} ({env['implementacion']})"],
            ["Algoritmo genético", f"DEAP {env['deap']} (creator, base.Toolbox, tools.selTournament, selRoulette, selBest, cxUniform, cxOnePoint)"],
            ["Cálculo y datos", f"NumPy {env['numpy']}, pandas {env['pandas']}"],
            ["Gráficos", f"Matplotlib {env['matplotlib']} (generados a partir de los CSV)"],
            ["Pruebas y notebook", f"pytest {env['pytest']}, nbconvert {env['nbconvert']} (notebook ejecutado de principio a fin)"],
            ["Informe", f"python-docx {env['python_docx']} sobre la plantilla oficial; exportación a PDF con Microsoft Word"],
            ["Verificación de legalidad", f"Node.js {c.legal['node'] if c.legal else '—'} + validador de Pokémon Showdown (opcional)"],
            ["Entorno de ejecución", f"{env['sistema']}, {env['procesador']}, {env['nucleos_logicos']} núcleos lógicos"]]
    R.table("herramientas", "Herramientas y versiones efectivamente utilizadas (results/environment.json).",
            ["Uso", "Herramienta y versión"], rows, [3.4, 8.8], ["l", "l"])
    R.h2("2.10. Asistentes de codificación basados en IA")
    R.p("En cumplimiento de la consigna se declara el uso de asistentes de IA. **ChatGPT** (OpenAI) se utilizó para preparar el "
        "prompt de trabajo que especifica los requisitos de esta entrega. **Claude Code** (Anthropic, modelo Claude Opus 5.5) se "
        "utilizó como asistente de desarrollo: a partir de la consigna, la propuesta aprobada, las presentaciones y los ejemplos "
        "de la cátedra, implementó el código, ejecutó los experimentos, generó los gráficos y tablas y redactó un borrador de este "
        "informe y de la guía de defensa. Los resultados provienen de ejecuciones reales registradas en el repositorio. La "
        "revisión final, la validación de los resultados y la defensa del trabajo corresponden a los integrantes del grupo.")

    # =========================================================================================
    R.h1("3. Resultados")
    R.p(f"Las {n_runs} corridas terminaron por el criterio de paro (200 generaciones) con 0 individuos inválidos registrados. "
        f"Cada corrida tomó en promedio {num(mean_time, 2)} s y realizó {ent(round(mean_req))} evaluaciones solicitadas, de las "
        f"cuales solo {ent(round(mean_eff))} fueron efectivas (combinaciones distintas calculadas; el resto se resolvió con la "
        f"caché), es decir, alrededor del {pct(mean_eff / CM4, 2)} del espacio.")
    R.h2("3.1. Corridas principales (semilla 42)")
    R.p("Las cinco corridas usan el **mismo cromosoma** (cuatro IDs distintos), la **misma función F** con los pesos de su perfil "
        "y los **mismos parámetros del AG** (@tab:params); solo cambia el vector de pesos. Cada gráfico muestra la aptitud "
        "máxima, media (con ±1 desvío) y mínima por generación, la mejor acumulada y, en el panel inferior, la cantidad de "
        "combinaciones distintas en la población. Por el elitismo, la máxima de cada generación coincide con la mejor acumulada.")
    interp = interpretations(c)
    for cid in ("R1", "R2", "R3", "R4", "R5"):
        s = c.main[cid]
        b = s["mejor_solucion"]
        w = prof[cid]
        log = c.main_logs[cid]
        opt = float(c.opt[cid]["optimo"])
        hit = abs(opt - b["aptitud"]) < 1e-9
        R.h3(f"3.1.{cid[1]}. {cid} — perfil «{pinfo[cid]['nombre']}»")
        R.p(f"**Función y parámetros.** F con w_{{C}} = {num(w.wC)}, w_{{P}} = {num(w.wP)}, w_{{A}} = {num(w.wA)}, w_{{S}} = "
            f"{num(w.wS)}, w_{{D}} = {num(w.wD)}; torneo k = 3, cruza uniforme (0,80; 0,5), mutación 0,10, elitismo 2, "
            f"100 individuos, 200 generaciones, semilla 42. La generación 0 tuvo máxima {num(float(log[0]['aptitud_max']), 3)} "
            f"y media {num(float(log[0]['aptitud_media']), 3)}; la mejor solución apareció en la **generación "
            f"{b['generacion_encontrada']}** (@fig:evol_{cid}). Evaluaciones: {ent(s['evaluaciones_solicitadas'])} solicitadas, "
            f"{ent(s['evaluaciones_efectivas'])} efectivas; tiempo {num(s['tiempo_segundos'], 2)} s.")
        R.figure(f"evol_{cid}", FIGURES_DIR / "runs" / f"{cid}_seed42.png",
                 f"Evolución de la aptitud en la corrida principal {cid} (perfil {pinfo[cid]['nombre'].lower()}, semilla 42). "
                 f"{interp[cid]['fig']}", 10.6)
        rows = [[f"{m['nombre_es']} ({cat.moves[m['move_id']].nombre_en})", m["tipo_es"], str(m["potencia"]), str(m["precision"]),
                 "Física" if m["categoria"] == "physical" else "Especial", "Sí" if m["stab"] else "No"] for m in b["movimientos"]]
        R.table(f"sol_{cid}", f"Mejor solución de la corrida {cid} (semilla 42).",
                ["Movimiento", "Tipo", "Potencia", "Precisión", "Categoría", "STAB"], rows, [4.4, 1.6, 1.5, 1.5, 1.7, 1.2],
                ["l", "l", "c", "c", "l", "c"])
        comp, aport = b["componentes"], b["aportes"]
        rows = [["Valor normalizado"] + [num(comp[k], 4) for k in "CPASD"] + ["—"],
                ["Peso"] + [num(getattr(w, "w" + k), 2) for k in "CPASD"] + ["Σ = 1"],
                ["Aporte (puntos de F)"] + [num(aport[k], 3) for k in "CPASD"] + [f"**{num(b['aptitud'], 4)}**"]]
        R.table(f"comp_{cid}", f"Componentes, pesos y aportes de la mejor solución de {cid}. Aptitud bajo el perfil base: "
                f"{num(b['aptitud_perfil_base'], 4)}.", ["", "C", "P", "A", "S", "D", "F"], rows,
                [3.2, 1.45, 1.45, 1.45, 1.45, 1.45, 1.75], ["l", "c", "c", "c", "c", "c", "c"], bold_first_col=True)
        R.p(f"**Cobertura:** {b['n_tipos_cubiertos']} de 18 tipos ({c.types_es(b['tipos_cubiertos'])}). **Sin cobertura "
            f"superefectiva:** {c.types_es(b['tipos_no_cubiertos'])}.")
        R.p(f"**Interpretación.** {interp[cid]['texto']}")
        n_ties = int(c.opt[cid]["n_empates_optimo"])
        R.p(f"**¿Resuelve el problema?** La solución es **válida** (cuatro movimientos distintos del catálogo; además, el validador "
            f"de Showdown confirma que pueden coexistir) y es la **mejor solución encontrada** en la corrida. "
            + (f"Su aptitud coincide con el **óptimo demostrado** por enumeración exhaustiva ({num(opt, 4)}"
               + (f"; hay {n_ties} combinaciones empatadas en el óptimo" if n_ties > 1 else "; el óptimo es único") + "), "
               "por lo que la corrida cumple el criterio de calidad y resuelve el problema **bajo el modelo adoptado**. "
               if hit else f"No alcanza el óptimo demostrado ({num(opt, 4)}). ")
            + interp[cid]["evaluacion"])

    R.h2("3.2. Comparación entre perfiles de pesos")
    rows = []
    for cid in ("R1", "R2", "R3", "R4", "R5"):
        b = c.main[cid]["mejor_solucion"]
        rows.append([cid, ", ".join(m["nombre_es"] for m in b["movimientos"])] + [num(b["componentes"][k], 3) for k in "CPASD"] +
                    [str(b["n_tipos_cubiertos"]), num(b["aptitud"], 3), num(b["aptitud_perfil_base"], 3)])
    R.table("perfiles_res", "Mejores soluciones de las corridas principales. F propia: con los pesos del perfil; F base: la misma "
            "solución evaluada con los pesos de R1 (referencia común).",
            ["Conf.", "Moveset", "C", "P", "A", "S", "D", "Tipos", "F propia", "F base"], rows,
            [0.9, 4.1, 0.85, 0.85, 0.85, 0.75, 0.6, 0.85, 1.15, 1.15], ["c", "l", "c", "c", "c", "c", "c", "c", "c", "c"], font_pt=7.5)
    same = [cid for cid in ("R1", "R2", "R4") if c.main[cid]["mejor_solucion"]["tipos_cubiertos"] == r1["tipos_cubiertos"] and
            "|".join(m["slug"] for m in c.main[cid]["mejor_solucion"]["movimientos"]) == r1_slugs]
    claim(same == ["R1", "R2", "R4"], "R1, R2 y R4 convergen al mismo moveset")
    claim(r1["n_tipos_cubiertos"] == max_cov, "el moveset de R1 alcanza la cobertura maxima")
    R.p(f"**R1, R2 y R4 convergen al mismo moveset** ({c.moveset_es(r1_slugs)}). La razón surge de la enumeración exhaustiva: "
        f"la cobertura máxima alcanzable con cuatro movimientos de este catálogo es **{max_cov} de 18 tipos** y la logran "
        f"{ent(n_max_cov)} combinaciones; entre ellas, ese moveset tiene la mejor combinación de potencia (P = {num(r1['componentes']['P'], 3)}) "
        f"y precisión (A = {num(r1['componentes']['A'], 4)}) y usa cuatro tipos distintos (D = 1). Aumentar el peso de la "
        "cobertura (R2) no puede mejorar C, y aumentar el de la precisión (R4) no compensa: las combinaciones de cobertura "
        f"máxima con precisión perfecta tienen menos potencia. Por eso el argmax no cambia, aunque el valor de F sí (@tab:perfiles_res).")
    R.p(f"**R3** desplaza la solución hacia tres movimientos de potencia 150, con P = 0,95 pero solo 9 tipos cubiertos; bajo el "
        f"perfil base vale {num(c.main['R3']['mejor_solucion']['aptitud_perfil_base'], 3)}, la más baja de las cinco. **R5** "
        f"incorpora un segundo movimiento con STAB (Anillo Ígneo) y obtiene S = 0,5 con 11 tipos cubiertos; bajo el perfil base "
        f"vale {num(c.main['R5']['mejor_solucion']['aptitud_perfil_base'], 4)}, apenas por debajo del óptimo base "
        f"({num(opt_r1, 4)}): es la segunda mejor combinación de todo el espacio para R1. Los valores de F de perfiles "
        "distintos no son comparables entre sí; por eso la comparación se hace con los componentes (@fig:componentes) y con la "
        "aptitud de referencia calculada con los pesos base (@tab:cruzada).")
    R.figure("componentes", FIGURES_DIR / "fig_componentes_perfiles.png",
             "Componentes normalizados de las soluciones óptimas de cada perfil. R1, R2 y R4 comparten solución; R3 gana potencia "
             "a costa de cobertura; R5 gana STAB a costa de cobertura y de un poco de precisión.", 11.6)
    rows = []
    for cid in ("R1", "R3", "R5"):
        x = c.cross[cid]
        rows.append([f"Solución de {cid}" + (" (= R2 = R4)" if cid == "R1" else "")] +
                    [num(float(x[f"aptitud_{p}"]), 3) for p in ("R1", "R2", "R3", "R4", "R5")])
    rows.append(["**Óptimo del perfil**"] + [f"**{num(float(c.opt[p]['optimo']), 3)}**" for p in ("R1", "R2", "R3", "R4", "R5")])
    R.table("cruzada", "Aptitud de cada solución principal evaluada con los cinco perfiles (filas: solución; columnas: perfil de "
            "evaluación).", ["Solución", "Pesos R1", "Pesos R2", "Pesos R3", "Pesos R4", "Pesos R5"], rows,
            [3.7, 1.7, 1.7, 1.7, 1.7, 1.7], ["l", "c", "c", "c", "c", "c"])

    R.h2("3.3. Réplicas y variabilidad")
    rows = []
    for e in c.exps["configs"]:
        a = c.configs[e.config_id]
        qa = c.qa[e.config_id]
        rows.append([e.config_id, num(float(a["aptitud_media"]), 3), num(float(a["aptitud_desvio"]), 3),
                     f"{num(float(a['aptitud_min']), 3)} – {num(float(a['aptitud_max']), 3)}", f"{qa['replicas_alcanzan_optimo']}/10",
                     f"{num(float(a['generacion_encontrada_media']), 1)} ({a['generacion_encontrada_min']}–{a['generacion_encontrada_max']})",
                     ent(round(float(a["evaluaciones_efectivas_media"]))), num(float(a["distintas_poblacion_final_media"]), 1),
                     a["movesets_distintos"]])
    R.table("replicas", "Resumen de las 10 réplicas por configuración (semillas 42–51). Desvío muestral. «Gen.»: generación en "
            "que apareció la mejor solución (media y rango). «Distintas final»: combinaciones distintas en la última población.",
            ["Conf.", "F media", "Desvío", "Rango", "Óptimo", "Gen. media (rango)", "Efectivas", "Distintas final", "Movesets"],
            rows, [1.5, 1.25, 1.0, 2.2, 1.0, 1.9, 1.15, 1.2, 1.1], ["l", "c", "c", "c", "c", "c", "c", "c", "c"], font_pt=7.5)
    claim(all(float(c.configs[e.config_id]["aptitud_desvio"]) == 0.0 for e in c.exps["configs"]), "desvio cero en todas")
    R.p(f"En las ocho configuraciones las diez réplicas terminaron con **la misma aptitud** (desvío 0): todas alcanzaron el "
        f"óptimo demostrado de su perfil ({hits_total} de {n_runs} corridas). En R3 las réplicas se reparten entre las **dos "
        "combinaciones empatadas** (Hiperrayo o Gigaimpacto, ambos Normal 150/90, con componentes idénticos); el modelo no "
        "puede distinguirlas. La variabilidad entre semillas aparece en el **costo** para llegar al óptimo: la generación de "
        f"hallazgo va de {gen_lo} a {gen_hi} en las configuraciones principales (@fig:variabilidad). Con diez réplicas, estas diferencias "
        "describen el comportamiento observado y no se atribuyen estadísticamente a los parámetros.")
    R.figure("variabilidad", FIGURES_DIR / "fig_variabilidad.png",
             "Variabilidad entre réplicas: generación en que cada réplica alcanzó el óptimo (izquierda) y evaluaciones efectivas "
             "acumuladas hasta ese momento (derecha, escala logarítmica). Cada punto es una semilla; la barra negra es la media.", 12.0)
    R.p("La aptitud media de la población **no aumenta en forma monótona**: en la corrida R1 disminuye respecto de la generación "
        f"anterior en {sum(1 for a, b in zip(c.main_logs['R1'], c.main_logs['R1'][1:]) if float(b['aptitud_media']) < float(a['aptitud_media']))} "
        "de las 200 transiciones, porque la mutación y la cruza siguen produciendo hijos peores aunque las élites preserven al "
        "mejor. A la vez, la diversidad cae rápido: en las configuraciones con torneo la población final contiene en promedio "
        "unas 10 combinaciones distintas de 100 individuos.")

    R.h2("3.4. Comparación de operadores")
    a = {k: c.configs[k] for k in ("R1", "OP-RUL", "OP-1P", "OP-MUT20")}
    R.p(f"Las tres variantes usan los pesos base y cambian una sola decisión respecto de R1 (@fig:operadores; figuras de semilla 42 en "
        f"el Anexo D). Todas alcanzaron el óptimo en 10/10 réplicas. La **ruleta** convergió más lento (generación media "
        f"{num(float(a['OP-RUL']['generacion_encontrada_media']), 1)} frente a {num(float(a['R1']['generacion_encontrada_media']), 1)}), "
        f"usó {ent(round(float(a['OP-RUL']['evaluaciones_efectivas_media'])))} evaluaciones efectivas en promedio (frente a "
        f"{ent(round(float(a['R1']['evaluaciones_efectivas_media'])))}) y mantuvo mucho más diversidad "
        f"({num(float(a['OP-RUL']['distintas_poblacion_final_media']), 1)} combinaciones distintas al final frente a "
        f"{num(float(a['R1']['distintas_poblacion_final_media']), 1)}): como todas las aptitudes de la población son del mismo "
        f"orden (en OP-RUL, semilla 42, la mejor nunca supera {num(rul_ratio, 1)} veces a la peor), las porciones de la ruleta "
        "son parecidas y la presión selectiva es baja. La **cruza de un punto** "
        f"se comportó como la uniforme (generación media {num(float(a['OP-1P']['generacion_encontrada_media']), 1)}). La "
        f"**mutación 0,20** duplicó aproximadamente la diversidad final ({num(float(a['OP-MUT20']['distintas_poblacion_final_media']), 1)}) "
        f"y bajó la aptitud media de la población, con una generación media de hallazgo de "
        f"{num(float(a['OP-MUT20']['generacion_encontrada_media']), 1)}.")
    R.figure("operadores", FIGURES_DIR / "fig_operadores.png",
             "Comparación de operadores con pesos base: mejor acumulada media (izquierda, primeras 120 generaciones) y aptitud media "
             "de la población (derecha), promediadas sobre 10 semillas.", 12.0)

    R.h2("3.5. Comparación con otros métodos")
    rows = []
    label = {"AG (10 semillas)": "AG (10 semillas)", "Busqueda aleatoria (mismas solicitadas)": "Aleatoria (= solicitadas)",
             "Busqueda aleatoria (mismas efectivas)": "Aleatoria (= efectivas)", "Voraz": "Voraz",
             "Voraz + busqueda local": "Voraz + búsqueda local", "Enumeracion exhaustiva": "Exhaustiva"}
    for m in c.methods:
        rows.append([m["perfil"], label[m["metodo"]],
                     num(float(m["aptitud_media"]), 3) + ("" if m["aptitud_min"] == m["aptitud_max"] else
                                                          f" ({num(float(m['aptitud_min']), 2)}–{num(float(m['aptitud_max']), 2)})"),
                     f"{m['replicas_alcanzan_optimo']}/{m['replicas']}", ent(round(float(m["evaluaciones_solicitadas_media"]))),
                     ent(round(float(m["evaluaciones_efectivas_media"]))), num(float(m["segundos_medios"]), 3)])
    R.table("metodos", "Comparación de métodos con el mismo catálogo y la misma función. Valores medios por perfil; entre "
            "paréntesis, rango de las 10 semillas. «Óptimo»: ejecuciones que alcanzan el óptimo demostrado.",
            ["Perfil", "Método", "Mejor F (rango)", "Óptimo", "Eval. solicitadas", "Eval. efectivas", "Tiempo (s)"],
            rows, [1.0, 3.0, 2.4, 1.1, 1.6, 1.5, 1.3], ["c", "l", "c", "c", "c", "c", "c"], font_pt=7)
    rs_req = [r for r in c.rs if r["modo_presupuesto"] == "requested"]
    rs_eff = [r for r in c.rs if r["modo_presupuesto"] == "effective"]
    claim(not any(r["alcanza_optimo"] == "True" for r in c.rs), "la aleatoria nunca alcanza el optimo")
    g_hits = sum(1 for g in c.greedy if g["metodo"] == "voraz" and g["alcanza_optimo"] == "True")
    ls_hits = sum(1 for g in c.greedy if g["metodo"] != "voraz" and g["alcanza_optimo"] == "True")
    best_rank_req = min(int(r["puesto_en_exhaustiva"]) for r in rs_req)
    R.p(f"**Búsqueda aleatoria.** Con el mismo número de evaluaciones solicitadas que el AG (unas {ent(round(mean_req))}, que "
        f"para la aleatoria son casi todas combinaciones distintas) no alcanzó el óptimo en ninguna de las 50 ejecuciones; su mejor "
        f"resultado quedó, en el mejor caso, en el puesto {best_rank_req} del ranking exhaustivo. Con el mismo número de "
        f"evaluaciones efectivas que el AG (~{ent(round(mean_eff))}) quedó todavía más lejos (@fig:ag_vs_aleatoria). El AG "
        "concentra la búsqueda en la región de alta aptitud: con muchas menos combinaciones distintas llega al óptimo.")
    R.figure("ag_vs_aleatoria", FIGURES_DIR / "fig_ag_vs_aleatoria.png",
             "Perfil base: mejor aptitud acumulada en función de las evaluaciones efectivas (combinaciones distintas), para las 10 "
             "semillas del AG y de la búsqueda aleatoria con presupuesto de evaluaciones solicitadas igual al del AG.", 12.0)
    gr = {(g["perfil"], g["metodo"]): g for g in c.greedy}
    g_miss = [p for p in ("R1", "R2", "R3", "R4", "R5") if gr[(p, "voraz")]["alcanza_optimo"] != "True"]
    claim(g_miss == ["R1", "R4"] and all(gr[(p, "voraz")]["puesto_en_exhaustiva"] == "2" for p in g_miss), "voraz segunda en R1 y R4")
    claim(gr[("R1", "voraz")]["detalle"].startswith("blast-burn"), "la voraz empieza por Anillo Igneo en R1")
    claim("blast-burn" not in r1_slugs, "el optimo R1 no incluye Anillo Igneo")
    ls_eff_max = max(int(g["evaluaciones_efectivas"]) for g in c.greedy if g["metodo"] != "voraz")
    claim(ls_eff_max < 600, "voraz+BL con menos de 600 efectivas")
    R.p(f"**Heurística voraz.** Con {ent(int(gr[('R1', 'voraz')]['evaluaciones_solicitadas']))} evaluaciones alcanzó el óptimo en "
        f"{g_hits} de 5 perfiles; en R1 y R4 quedó en el segundo puesto del ranking exhaustivo. En R1 eligió primero Anillo Ígneo, "
        "que tiene la mayor ganancia marginal inicial (STAB, potencia 150 y cobertura de cuatro tipos), y la construcción ya no "
        "pudo llegar al óptimo, que no incluye ese movimiento: terminó en la solución de R5. **Voraz + búsqueda local** alcanzó "
        f"el óptimo en {ls_hits} de 5 perfiles con a lo sumo {ent(ls_eff_max)} evaluaciones efectivas: en un espacio de este "
        "tamaño y con esta función, una heurística simple es competitiva.")
    R.p(f"**Enumeración exhaustiva.** Evaluó las {ent(CM4)} combinaciones en {num(ex_secs, 2)} s (cálculo de componentes, que "
        "no dependen de los pesos) más unos milisegundos por perfil, y entrega el **óptimo demostrado** y la cantidad de empates. "
        "Para este catálogo es el método más directo y confiable; la @fig:distribucion muestra lo excepcional del óptimo dentro de "
        "la distribución de aptitudes.")
    R.figure("distribucion", FIGURES_DIR / "fig_distribucion_exhaustiva.png",
             "Distribución de la aptitud (perfil base) de las 292.825 combinaciones posibles, con el óptimo exacto (alcanzado por el "
             "AG en todas las semillas), el resultado de la heurística voraz y la media de la búsqueda aleatoria.", 12.0)

    R.h2("3.6. Sensibilidad del catálogo: subconjunto conservador")
    rows = []
    for pid in ("R1", "R2", "R3", "R4", "R5"):
        x = c.cons[pid]
        ms = x["movesets_optimos_es"].split(" || ")
        rows.append([pid, ms[0].replace(" | ", ", ") + (f" (empate: {len(ms)})" if len(ms) > 1 else ""),
                     x["n_tipos_cubiertos"], num(float(x["optimo_restringido"]), 3), num(float(x["optimo_catalogo_completo"]), 3),
                     ent(int(x["puesto_en_catalogo_completo"]))])
    R.table("conservador", f"Óptimo exacto restringido al subconjunto conservador ({counts['subconjunto_conservador_M']} movimientos, "
            f"{ent(int(c.ex_info['combinaciones_conservador']))} combinaciones), con la misma función y normalización.",
            ["Perfil", "Moveset óptimo restringido", "Tipos", "F restringida", "F catálogo completo", "Puesto en el completo"],
            rows, [1.0, 5.3, 0.9, 1.5, 1.6, 1.7], ["c", "l", "c", "c", "c", "c"], font_pt=7.5)
    R.p("Las soluciones del catálogo completo dependen de movimientos con costos que el modelo ignora: en los cinco perfiles "
        "aparece Puño Certero y, según el perfil, Vuelo, Rayo Solar, Anillo Ígneo, Hiperrayo o Gigaimpacto. Si se excluyen por "
        "regla los movimientos con carga, recarga o fallo condicional, el óptimo cambia a combinaciones como "
        f"{c.cons['R1']['movesets_optimos_es'].replace(' | ', ', ')} para el perfil base, con menor aptitud bajo el modelo. "
        "Este análisis no modifica la función aprobada: muestra cuánto depende el resultado de la definición del catálogo.")

    R.h2("3.7. Verificación de legalidad de los movesets")
    if c.legal:
        rows = []
        for r in c.legal_rows:
            if r["tipo_caso"] == "moveset_resultante":
                rows.append([c.moveset_es(r["moveset"]), r["origen"].replace("busqueda_local", "búsqueda local").replace("optimo", "óptimo"),
                             "Válido" if r["valido_showdown"] == "True" else "Inválido"])
        for r in c.legal_rows:
            if r["tipo_caso"] != "moveset_resultante":
                rows.append([r["movimientos_en"], "Control negativo: " + r["origen"].replace("Thunderbolt", "Thunderbolt (Rayo)"),
                             "Rechazado: " + r["mensaje"]])
        claim(meta["estadisticas_base_charizard"]["special-attack"] > meta["estadisticas_base_charizard"]["attack"], "SpA > Atk")
        R.table("legalidad", f"Validación con Pokémon Showdown («[Gen 9] Anything Goes», regla Obtainable, commit {sd_commit[:7]}).",
                ["Moveset", "Origen", "Resultado"], rows, [4.6, 5.0, 2.6], ["l", "l", "l"], font_pt=7)
    R.h2("3.8. Validación del software")
    if c.validation:
        vt = c.validation["pytest"]
        R.p(f"La suite de pruebas (`tests/`, {vt['total']} pruebas) se ejecutó con resultado **{vt['passed']} aprobadas y "
            f"{vt['failed']} fallidas**. Cubre: fórmulas de C/P/A/S/D y F contra un cálculo manual independiente; aptitud 0 ante "
            "longitud incorrecta, repetidos o IDs no permitidos; invariancia ante las 24 permutaciones; orientación y valores de "
            "la tabla de tipos (incluidas las inmunidades); inicialización, cruza con reparación y mutación válidas; tamaño "
            "constante, élites y mejor acumulada coherentes; reproducibilidad con la misma semilla; recreación de clases de DEAP "
            "sin advertencias; enumeración exhaustiva frente a fuerza bruta; y coherencia de logs, resúmenes y gráficos. "
            + " ".join(c.validation.get("frases_informe", [])))
    else:
        R.p("La validación del software se registra en docs/registro_validacion.md.")

    # =========================================================================================
    R.h1("4. Discusión")
    R.h2("4.1. ¿El sistema resuelve satisfactoriamente el problema?")
    R.p(f"**Bajo el modelo aprobado, sí.** Las {n_runs} corridas produjeron soluciones válidas y alcanzaron el óptimo demostrado "
        "de su perfil; las ocho configuraciones cumplen el criterio de calidad fijado de antemano (10/10 réplicas en el óptimo, "
        "brecha 0). Que el AG lo logre en todas las semillas **no prueba que un AG garantice el óptimo**: la garantía proviene "
        "de la enumeración exhaustiva, y la teoría advierte que la población final no asegura el óptimo global (teoría, pág. 43).")
    R.p("**Como recomendación competitiva real, solo parcialmente.** La aptitud es un indicador simplificado de calidad ofensiva: "
        "no equivale a probabilidad de victoria ni demuestra cuál es el mejor Charizard competitivo. El AG optimizó exactamente "
        "lo que la función premia, y la función ignora costos importantes: Vuelo y Rayo Solar necesitan un turno de carga (Rayo "
        "Solar no lo necesita con sol, condición fuera del alcance), Puño Certero falla si Charizard recibe daño antes de actuar y Anillo Ígneo, "
        "Hiperrayo y Gigaimpacto obligan a recargar. Otras simplificaciones relevantes:")
    R.bullets([
        "**Tipos individuales**: la cobertura se mide contra 18 tipos y no contra combinaciones de doble tipo; un ataque "
        "superefectivo contra un tipo puede ser neutro o resistido contra un Pokémon de doble tipo.",
        f"**Reparto físico/especial**: Charizard tiene más Ataque Especial que Ataque ({meta['estadisticas_base_charizard']['special-attack']} "
        f"frente a {meta['estadisticas_base_charizard']['attack']} de estadística base según PokéAPI), pero el modelo trata igual un "
        "movimiento físico y uno especial; varias soluciones mezclan ambas categorías.",
        "**Dependencias entre criterios**: P y A se promedian por separado, de modo que un movimiento potente e impreciso y "
        "uno débil y preciso pueden compensarse; la cobertura no penaliza inmunidades ni resistencias.",
        "**STAB** se cuenta como presencia (S) y no multiplica la potencia, tal como se aprobó; por eso el modelo no prefiere "
        "naturalmente los ataques Fuego/Volador.",
    ])
    R.p("Estas limitaciones se documentan en lugar de cambiar silenciosamente la función aprobada. El análisis con el subconjunto "
        "conservador (Sección 3.6) muestra el efecto de la principal de ellas.")
    R.h2("4.2. Comportamiento esperado y observado")
    R.p("La teoría anticipa que el torneo es simple pero que puede eliminar la diversidad rápidamente y llevar a estancamiento "
        "(teoría, pág. 24), que la ruleta mantiene diversidad pero es sensible a la escala de las aptitudes (págs. 27–28), que "
        "una mutación alta aumenta la exploración a costa de la herencia (pág. 40) y que la convergencia prematura es un riesgo "
        "(pág. 26). Lo observado es coherente: con torneo la población pasó de 100 a unas 10 combinaciones distintas en pocas "
        f"decenas de generaciones; la ruleta mantuvo unas {num(div['OP-RUL'] / div['R1'], 1)} veces más diversidad y convergió más "
        f"lento; la mutación 0,20 multiplicó por {num(div['OP-MUT20'] / div['R1'], 1)} la diversidad final. En este problema la "
        "pérdida de diversidad no impidió llegar al óptimo, porque el óptimo se alcanza temprano (generación media entre "
        f"{num(min(gmeans), 1)} y {num(max(gmeans), 1)} en las configuraciones principales) y la superficie de aptitud tiene pocos "
        "óptimos locales relevantes, como sugiere el éxito de la búsqueda local. La aptitud media no creció en forma monótona.")
    R.h2("4.3. ¿Cómo se compara con otros métodos?")
    R.p(f"Frente a la **búsqueda aleatoria**, el AG fue claramente superior con presupuesto igualado: llegó al óptimo en todas "
        f"las semillas evaluando unas {ent(round(mean_eff))} combinaciones distintas, mientras que la aleatoria no lo alcanzó "
        f"nunca ni con unas {ent(round(statistics.fmean(int(r['evaluaciones_efectivas']) for r in rs_req)))} combinaciones distintas. "
        f"Frente a la **heurística voraz**, el AG fue más robusto ({g_hits}/5 perfiles para la voraz). Sin embargo, **la "
        f"enumeración exhaustiva es la mejor opción para este tamaño de problema**: en {num(ex_secs, 2)} s entrega el óptimo exacto "
        "de todos los perfiles y sus empates, sin parámetros ni semillas, y voraz + búsqueda local también llegó al óptimo en los "
        "cinco perfiles. Este es un resultado válido del trabajo: el beneficio del AG aquí es metodológico (implementar, "
        "parametrizar y analizar un proceso evolutivo, como pide la asignatura) y de escalabilidad, no de necesidad. El AG "
        "aportaría más en variantes donde la enumeración deja de ser viable o la aptitud no se puede descomponer: por ejemplo, "
        "optimizar a la vez los movimientos de los seis Pokémon de un equipo (el espacio crece a C(M,4)^6) o evaluar con "
        "simulaciones de combate costosas.")
    R.h2("4.4. ¿Qué se podría mejorar con más tiempo y recursos?")
    R.bullets([
        "**Modelo de aptitud**: potencia efectiva por turno (penalizar carga/recarga y fallo condicional), consistencia "
        "física/especial según las estadísticas de Charizard, cobertura ponderada por la frecuencia de tipos (incluidos dobles) "
        "en el metajuego y penalización de inmunidades.",
        "**Evaluación por simulación**: estimar daño esperado contra un conjunto representativo de rivales con un simulador "
        "(por ejemplo, el motor de Showdown) en lugar de componentes estáticos.",
        "**Optimización multiobjetivo**: obtener el frente de Pareto (por ejemplo con NSGA-II, disponible en DEAP) en lugar de "
        "fijar pesos, y mostrar al jugador las alternativas no dominadas.",
        "**Diversidad**: nichos o *crowding*, mutación adaptativa (teoría, pág. 39) y paro por estancamiento registrado como "
        "criterio adicional; análisis estadístico con más réplicas.",
        "**Alcance**: incorporar habilidades (Mar Llamas, Poder Solar), objetos, clima y teracristalización, y extender la "
        "optimización al equipo completo, donde un AG sí sería necesario.",
    ])

    # =========================================================================================
    R.h1("5. Conclusión")
    R.h2("5.1. Conclusiones de la implementación")
    R.p("Se implementó un AG completo, reproducible y verificado con DEAP para optimizar el moveset de Charizard con la función "
        "aprobada, sobre datos reales y trazables de Pokémon Escarlata/Púrpura. Las cinco corridas principales y sus 75 réplicas "
        "adicionales encontraron el óptimo demostrado de cada perfil; la cobertura máxima posible es de 13 tipos y el perfil base "
        f"la alcanza con {c.moveset_es(r1_slugs)}. Los pesos determinan qué se optimiza: tres perfiles comparten solución porque "
        "esta domina en cobertura, precisión y diversidad, mientras que dar más peso a la potencia o al STAB cambia el moveset.")
    R.h2("5.2. Relación con la teoría vista en clase")
    R.p("Los elementos del AG se corresponden con la estructura vista en clase: individuos y cromosoma (págs. 13–17), función de "
        "aptitud con penalización dura (págs. 18–21), población inicial aleatoria (pág. 22), selección por torneo y ruleta "
        "(págs. 23–28), cruza simple y binomial aleatoria (págs. 33–37), mutación simple (págs. 38–40), criterio de paro por "
        "generaciones (pág. 41) y población final (págs. 42–43). Se confirmó en la práctica que la presión de selección controla "
        "el equilibrio entre explotación y diversidad, y que un AG encuentra buenas soluciones sin garantizar la mejor: para "
        "afirmarlo fue necesaria una referencia exacta.")
    R.h2("5.3. Problemas encontrados y estrategias de resolución")
    R.bullets([
        "**Duplicados generados por la cruza** (un hijo con el mismo movimiento dos veces): reparación que conserva genes válidos "
        "y completa con genes de los padres; el validador sigue asignando 0 como resguardo.",
        "**Redefinición de clases** con `creator.create` al reejecutar celdas: creación única con `ensure_deap_types` y un "
        "`Toolbox` nuevo por corrida.",
        "**Reproducibilidad**: una sola fuente de aleatoriedad por corrida (`random`, usada también por DEAP) inicializada con la "
        "semilla; prueba automática de determinismo.",
        "**Ruleta con suma de aptitudes 0**: caso contemplado con selección uniforme.",
        "**Datos**: semántica de `past_values` de PokéAPI para reconstruir atributos de la versión; precisión nula en movimientos "
        "que nunca fallan y un movimiento de estado con potencia 0 en la fuente, resueltos con reglas de exclusión explícitas.",
        "**Banderas mecánicas mal leídas**: al principio Vuelo y Excavar no quedaban marcados como movimientos con carga porque "
        "sus banderas ocupan varias líneas en Showdown; se corrigió el lector y se agregó la detección de movimientos de dos turnos.",
        "**Diferencia de learnset** entre fuentes (Falso Tortazo heredado de preevoluciones): se documentó el criterio.",
        "**Comparación numérica con la exhaustiva**: el óptimo calculado con producto matricial difería en 2·10^{-16} del valor "
        "del AG; se recalcula con la misma rutina de la aptitud y se usa una tolerancia de 10^{-9} para empates.",
        "**Legalidad conjunta**: el validador de Showdown exige EV no nulos; se agrega 1 EV solo para pasar ese control de formato.",
    ])
    R.h2("5.4. Lecciones aprendidas")
    R.bullets([
        "La definición de la función de aptitud y del catálogo pesa más que los operadores: el AG encontró exactamente lo que "
        "la función premia, incluidas debilidades del modelo.",
        "Medir una metaheurística exige una referencia: sin la enumeración exhaustiva no se podría distinguir «óptimo» de "
        "«mejor encontrado».",
        "Fijar versiones de datos, semillas y configuraciones, y generar el informe desde los logs, hace que cada número sea "
        "rastreable y que los resultados puedan reproducirse.",
        "Antes de elegir un método conviene medir el tamaño real del espacio: para C(53,4) una búsqueda exhaustiva es viable.",
    ])

    # =========================================================================================
    R.h1("6. Referencias")
    R.refs([
        "Cátedra de Inteligencia Artificial, UTN FRBA (2026). *Trabajo Práctico de Sistemas Inteligentes* (consigna TP2), "
        "2.º cuatrimestre de 2026. Campus Virtual.",
        "Grupo 6 (2026). *Trabajo Práctico N.º 2 – Entrega Parcial*: Propuesta 1, «Optimización del moveset ofensivo del Pokémon "
        "Charizard mediante Algoritmo Genético». Presentada el 28/09/2026.",
        "Pollo Cattaneo, M. F.; Castro, J. *Algoritmos Genéticos* [presentación de clase, 45 págs.]. Inteligencia Artificial, UTN FRBA. "
        "Págs. 12–43.",
        "Pollo Cattaneo, M. F. *Algoritmos Genéticos – Práctica* [presentación de clase, 26 págs.]. Inteligencia Artificial, UTN FRBA. "
        "Desarrollo propio (págs. 6–18) y DEAP (págs. 20–24).",
        "Cátedra de Inteligencia Artificial, UTN FRBA. *demoAlgoritmosGeneticos* (ZIP provisto por la cátedra): README.md, "
        "AG-OneMaxProblem.ipynb, AG-ViajeroProblem.ipynb y Ejemplo.ipynb.",
        "Holland, J. H. (1975). *Adaptation in Natural and Artificial Systems*. University of Michigan Press (citado en [3]).",
        "Fortin, F.-A.; De Rainville, F.-M.; Gardner, M.-A.; Parizeau, M.; Gagné, C. (2012). DEAP: Evolutionary Algorithms Made "
        "Easy. *Journal of Machine Learning Research*, 13, 2171–2175.",
        f"DEAP Project. *DEAP documentation* (versión {c.env['deap']}). https://deap.readthedocs.io/ (consulta: 03/10/2026).",
        f"PokéAPI. *The RESTful Pokémon API*. https://pokeapi.co/ ; datos: https://github.com/PokeAPI/api-data, commit {pk_commit} "
        f"(consulta: {fecha_datos.split('-')[2]}/{fecha_datos.split('-')[1]}/{fecha_datos.split('-')[0]}). Fuente comunitaria.",
        f"Smogon. *Pokémon Showdown*. https://github.com/smogon/pokemon-showdown, commit {sd_commit} (consulta: 03/10/2026). "
        "Fuente comunitaria.",
        "Harris, C. R. et al. (2020). Array programming with NumPy. *Nature*, 585, 357–362.",
        "Hunter, J. D. (2007). Matplotlib: A 2D graphics environment. *Computing in Science & Engineering*, 9(3), 90–95.",
        "The pandas development team. *pandas*. https://pandas.pydata.org/ (consulta: 03/10/2026).",
        "python-docx. https://python-docx.readthedocs.io/ (consulta: 03/10/2026).",
        f"Repositorio del proyecto (código fuente, datos y logs CSV de las corridas): {repo_txt}",
    ])

    # =========================================================================================
    R.page_break()
    R.h1("Anexo A. Matriz de cumplimiento")
    R.p("Relación entre los requisitos de la consigna aplicables a un AG, la sección del informe y la evidencia en el repositorio. "
        "La versión completa está en `docs/matriz_cumplimiento.md`. Los requisitos exclusivos de redes neuronales no corresponden.")
    R.table("matriz", "Matriz de cumplimiento (resumen).", ["Requisito", "Sección", "Evidencia", "Estado"],
            [[a, b_, d, e] for a, b_, d, e in compliance_rows(c, repo_url)], [4.3, 1.4, 4.6, 1.9], ["l", "c", "l", "l"], font_pt=7)
    R.h1("Anexo B. Catálogo definitivo")
    rows = []
    for r in sorted(c.catalog_rows, key=lambda r: (c.tes(r["tipo"]), r["nombre_es"])):
        rows.append([f"{r['nombre_es']} ({r['nombre_en']})", r["tipo_es"], r["potencia"], r["precision"],
                     "F" if r["categoria"] == "physical" else "E", "Sí" if r["stab"] == "True" else "",
                     r["metodos_sv"].replace("level-up", "nivel").replace("machine", "MT").replace("egg", "huevo"),
                     r["efectos_ignorados"] or "—"])
    R.table("catalogo", f"Catálogo definitivo (M = {M}). Cat.: F física, E especial. Métodos en Escarlata/Púrpura. Última columna: "
            "efectos que la función aprobada no modela.",
            ["Movimiento", "Tipo", "Pot.", "Prec.", "Cat.", "STAB", "Métodos", "Efectos ignorados"], rows,
            [3.1, 1.2, 0.7, 0.7, 0.6, 0.7, 1.9, 3.3], ["l", "l", "c", "c", "c", "c", "l", "l"], font_pt=6.5)
    R.h1("Anexo C. Movimientos excluidos")
    motivo = {"ESTADO": "Estado (sin daño directo)", "POTENCIA_NO_NUMERICA": "Potencia no numérica",
              "PRECISION_NO_NUMERICA": "Precisión no numérica (nunca falla)", "GOLPES_MULTIPLES": "Golpes múltiples",
              "POTENCIA_CERO": "Potencia 0 en la fuente", "OHKO": "KO en un golpe"}
    rows = [[f"{r['nombre_es']} ({r['nombre_en']})", r["tipo_es"], r["potencia"] or "—", r["precision"] or "—",
             {"status": "estado", "physical": "física", "special": "especial"}[r["clase_danio"]],
             motivo[r["motivo_exclusion"]] + (" (fuente: pot. 0, prec. 0)" if r["potencia"] == "0" else "")]
            for r in c.learnset if r["incluido"] == "False"]
    R.table("excluidos", f"Movimientos del learnset excluidos del catálogo ({len(rows)}).",
            ["Movimiento", "Tipo", "Pot.", "Prec.", "Clase", "Motivo"], rows, [3.8, 1.4, 0.9, 0.9, 1.4, 3.8],
            ["l", "l", "c", "c", "l", "l"], font_pt=7)
    R.h1("Anexo D. Evolución de las variantes de operadores (semilla 42)")
    for cid, desc in (("OP-RUL", "selección por ruleta"), ("OP-1P", "cruza de un punto"), ("OP-MUT20", "mutación 0,20")):
        R.figure(f"evol_{cid.replace('-', '_')}", FIGURES_DIR / "runs" / f"{cid}_seed42.png",
                 f"Evolución de la aptitud en {cid} ({desc}, pesos base, semilla 42).", 10.5)
    R.p("Los gráficos de las 80 corridas (8 configuraciones × 10 semillas) están en `figures/runs/` y sus logs en `results/runs/`.")
    R.h1("Anexo E. Reproducción")
    R.p("Comandos (Windows; en Linux/macOS reemplazar `.venv\\Scripts\\python` por `.venv/bin/python`). Detalle en README.md.")
    for cmd in ("python -m venv .venv", ".venv\\Scripts\\python -m pip install -r requirements.txt",
                ".venv\\Scripts\\python scripts\\run_ga.py --config R1 --seed 42",
                ".venv\\Scripts\\python scripts\\run_all.py"):
        R.p(f"`{cmd}`", size=8.5)
    return R


def interpretations(c: Ctx) -> dict:
    """Textos de interpretacion por corrida. Cada afirmacion se apoya en datos verificados con claim()."""
    m = {cid: c.main[cid]["mejor_solucion"] for cid in c.main}
    slugs = {cid: {x["slug"] for x in m[cid]["movimientos"]} for cid in m}
    claim(slugs["R1"] == {"fly", "solar-beam", "earthquake", "focus-punch"}, "moveset R1")
    claim(slugs["R3"] in ({"hyper-beam", "solar-beam", "focus-punch", "blast-burn"}, {"giga-impact", "solar-beam", "focus-punch", "blast-burn"}), "moveset R3")
    claim(slugs["R5"] == {"fly", "earthquake", "focus-punch", "blast-burn"}, "moveset R5")
    claim(int(c.opt["R3"]["n_empates_optimo"]) == 2, "empate en R3")
    claim(slugs["R2"] == slugs["R1"] == slugs["R4"], "R1 = R2 = R4")
    claim(m["R1"]["n_tipos_cubiertos"] == 13 and m["R3"]["n_tipos_cubiertos"] == 9 and m["R5"]["n_tipos_cubiertos"] == 11, "coberturas")
    claim(abs(m["R3"]["componentes"]["P"] - 0.95) < 1e-12 and abs(m["R5"]["componentes"]["S"] - 0.5) < 1e-12, "P de R3 y S de R5")
    top_r1 = [t for t in c.top if t["perfil"] == "R1"]
    claim(set(top_r1[1]["moveset"].split("|")) == slugs["R5"], "solucion R5 es la segunda para R1")
    g = {cid: m[cid]["generacion_encontrada"] for cid in m}
    return {
        "R1": {
            "fig": f"La mejor solución aparece en la generación {g['R1']} y la diversidad cae a unas 10 combinaciones distintas.",
            "texto": "El moveset combina cuatro tipos distintos (D = 1) y alcanza la cobertura máxima posible del catálogo, 13 tipos: "
                     "Puño Certero (Lucha) cubre Normal, Hielo, Roca, Siniestro y Acero; Terremoto (Tierra) agrega Fuego, Eléctrico y "
                     "Veneno; Rayo Solar (Planta) agrega Agua y Tierra; y Vuelo (Volador) agrega Lucha, Planta y Bicho. Solo Vuelo "
                     "recibe STAB (S = 0,25). La potencia y la precisión son altas porque el modelo valora la potencia base de Puño "
                     "Certero (150) y de Rayo Solar (120) sin descontar que el primero falla si Charizard recibe daño y que el "
                     "segundo, como Vuelo, necesita un turno de carga.",
            "evaluacion": "Desde el punto de vista competitivo la solución debe leerse con cautela por esos costos ignorados "
                          "(Sección 4.1); el subconjunto conservador ofrece una alternativa (Sección 3.6).",
        },
        "R2": {
            "fig": f"La mejor solución aparece en la generación {g['R2']}; el aumento de w_{{C}} eleva la escala de F respecto de R1.",
            "texto": "Con más peso en la cobertura el AG converge al **mismo moveset que R1**: ese moveset ya tiene la cobertura "
                     "máxima alcanzable (13/18), por lo que el aumento de w_{C} no puede mejorar C y la elección entre las 300 "
                     "combinaciones de cobertura máxima la deciden la potencia y la precisión. Lo que cambia es la composición de F: "
                     f"la cobertura aporta {num(m['R2']['aportes']['C'], 1)} de {num(m['R2']['aptitud'], 2)} puntos.",
            "evaluacion": "Las mismas salvedades de R1 aplican a la interpretación competitiva.",
        },
        "R3": {
            "fig": f"La mejor solución aparece en la generación {g['R3']}.",
            "texto": "Con w_{P} = 0,45 el AG elige tres movimientos de potencia 150 y Rayo Solar (P = 0,95), resignando cobertura: "
                     "solo 9 tipos cubiertos, sin respuesta superefectiva contra Fuego, Eléctrico, Lucha ni Veneno. Hiperrayo y "
                     "Gigaimpacto son intercambiables (ambos Normal, 150/90): el óptimo tiene **dos combinaciones empatadas** y la "
                     "semilla decide cuál aparece. Tres de los cuatro movimientos obligan a recargar o pueden fallar.",
            "evaluacion": "Es el perfil donde la simplificación del modelo más se nota: la potencia base no descuenta los turnos de recarga.",
        },
        "R4": {
            "fig": f"La mejor solución aparece en la generación {g['R4']}.",
            "texto": "Con w_{A} = 0,40 el AG vuelve a elegir el **mismo moveset que R1**. Ya tiene A = 0,9875 (solo Vuelo tiene 95 % "
                     "de precisión); las combinaciones de cobertura máxima con precisión perfecta reemplazan a Vuelo por "
                     "movimientos más débiles y pierden más en P y en F que lo que ganan en A. "
                     f"La precisión aporta {num(m['R4']['aportes']['A'], 1)} de {num(m['R4']['aptitud'], 2)} puntos.",
            "evaluacion": "Las mismas salvedades de R1 aplican a la interpretación competitiva.",
        },
        "R5": {
            "fig": f"La mejor solución aparece en la generación {g['R5']}.",
            "texto": "Con w_{S} = 0,30 y w_{D} = 0,20 el AG incorpora un segundo ataque con STAB, Anillo Ígneo (Fuego, 150/90), en lugar de "
                     "Rayo Solar: S sube a 0,5 y la diversidad se mantiene en 1, pero la cobertura baja a 11 tipos (se pierden Agua "
                     "y Tierra). Bajo el perfil base esta solución es la segunda mejor de todo el espacio.",
            "evaluacion": "Competitivamente aprovecha mejor los tipos de Charizard, pero depende de Puño Certero, Vuelo y de un "
                          "movimiento con recarga.",
        },
    }


def compliance_rows(c: Ctx, repo_url) -> list[tuple]:
    """Fuente unica de la matriz de cumplimiento (se usa en el Anexo A y en docs/matriz_cumplimiento.md)."""
    n = len(c.runs)
    legal = c.legal is not None
    val = c.validation is not None and c.validation["pytest"]["failed"] == 0
    return [
        ("Problema con relevancia; descripción, características y ejemplo", "1.1–1.3", "Informe; entrega parcial", "Cumplido"),
        ("No figura en el anexo de problemas resueltos", "1.4", "Consigna, anexo (pág. 25)", "Cumplido"),
        ("Justificación del AG (sin afirmar imposibilidad; tamaño real del espacio)", "1.6, 4.3", "C(53,4) medido; exhaustiva", "Cumplido"),
        ("Cromosoma y función de aptitud con justificación; versión por corrida", "2.4–2.6", "src/charizard_ga/fitness.py; logs (fitness_version)", "Cumplido"),
        ("Métodos de los operadores genéticos usados", "2.7", "src/charizard_ga/operators.py, ga.py", "Cumplido"),
        ("Herramientas, lenguaje y librerías", "2.9", "requirements.txt; results/environment.json", "Cumplido"),
        ("Asistente de codificación basado en IA", "2.10", "Declaración en el informe", "Cumplido"),
        ("Al menos 5 corridas con parámetros distintos", "2.8, 3.1", f"{n} corridas en results/runs/", "Cumplido"),
        ("Por corrida: cambios en F/cromosoma, métodos y parámetros", "3.1", "results/parameters.json; summaries", "Cumplido"),
        ("Por corrida: gráfico de toda la corrida", "3.1, Anexo D", "figures/runs/*.png (desde CSV)", "Cumplido"),
        ("Por corrida: mejor solución, interpretación y si resuelve", "3.1", "seed_42_summary.json; ga_vs_optimum.csv", "Cumplido"),
        ("Discusión: comparación, ¿resuelve?, otros métodos, mejoras", "3.2–3.6, 4", "results/baselines/", "Cumplido"),
        ("Conclusión: implementación, teoría, problemas, lecciones", "5", "Informe", "Cumplido"),
        ("Referencias bibliográficas, datos y fuentes", "6", "data/FUENTES_DATOS.md", "Cumplido"),
        ("URL del repositorio con código fuente", "6", repo_url or "Carpeta lista para publicar", "Cumplido" if repo_url else "Pendiente: URL"),
        ("Log CSV de las corridas principales en el repositorio", "6", "results/runs/R1..R5/seed_42_log.csv", "Cumplido (falta publicar)" if not repo_url else "Cumplido"),
        ("Informe con la plantilla oficial", "Todo", "docs/plantilla/ (plantilla del Campus)", "Cumplido"),
        ("Datos versionados y ejecución sin conexión", "2.1–2.2", "data/; scripts/prepare_data.py", "Cumplido"),
        ("Coexistencia de los movimientos", "2.2, 3.7", "results/legality_showdown.csv", "Cumplido" if legal else "Limitado"),
        ("Pruebas y registro de validación", "3.8", "tests/; docs/registro_validacion.md", "Cumplido" if val else "Revisar"),
        ("Notebook ejecutable de principio a fin", "—", "notebooks/tp2_charizard.ipynb", "Cumplido"),
    ]


# ---------------------------------------------------------------------------
# Plantilla: portada y encabezado del cuerpo
# ---------------------------------------------------------------------------
def fill_template(doc, c: Ctx, abstract: str, keywords: str):
    rc = c.report_cfg
    replace_text_everywhere(doc, "Año 20xx", f"Año {rc['anio']}")
    replace_text_everywhere(doc, "Trabajo Práctico N° xxx", rc["trabajo_practico"])
    replace_text_everywhere(doc, "GRUPO N° xx", f"GRUPO N° {rc['grupo']}")
    replace_text_everywhere(doc, "xx/xx/20xx", rc["fecha_presentacion"])
    members = rc["integrantes"]
    group_table = doc.tables[1]
    empty_rows = len(group_table.rows) - 3
    if len(members) > empty_rows:
        duplicate_last_row(group_table, len(members) - empty_rows)
    for i, (name, mail) in enumerate(members):
        row = group_table.rows[3 + i]
        set_cell_text(row.cells[0], name, size=10)
        set_cell_text(row.cells[1], mail, size=10)
    # Cuerpo: titulo, autores, resumen y palabras clave en los estilos de la plantilla.
    for p in doc.paragraphs:
        if p.style.name == "title":
            _set_par(p, rc["titulo"])
        elif p.style.name == "author":
            _set_par(p, ", ".join(_nice_name(n) for n, _ in members))
        elif p.style.name == "abstract" and p.text.startswith("Resumen"):
            _set_par(p, "")
            r = p.add_run("Resumen. ")
            r.bold = True
            p.add_run(abstract)
        elif p.style.name == "abstract" and p.text.startswith("Palabras"):
            _set_par(p, "")
            r = p.add_run("Palabras Claves: ")
            r.bold = True
            p.add_run(keywords)
    remove_body_from(doc, "1. Introducción")


def _nice_name(apellido_nombre: str) -> str:
    ap, nom = [x.strip() for x in apellido_nombre.split(",")]
    return f"{nom} {ap}"


def _set_par(p, text):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    if text:
        p.add_run(text)


def abstract_text(c: Ctx) -> str:
    r1 = c.main["R1"]["mejor_solucion"]
    return (f"Se implementó un Algoritmo Genético con DEAP para elegir los cuatro movimientos ofensivos de Charizard (Fuego/Volador) "
            f"en Pokémon Escarlata/Púrpura, maximizando una función ponderada de cobertura de tipos, potencia, precisión, STAB y "
            f"diversidad. El catálogo de {c.cat.M} movimientos se obtuvo de PokéAPI y se verificó contra Pokémon Showdown. Se "
            f"ejecutaron cinco perfiles de pesos y tres variantes de operadores, con diez semillas cada uno ({len(c.runs)} corridas). "
            f"La enumeración exhaustiva de las {ent(math.comb(c.cat.M, 4))} combinaciones confirmó que todas las corridas alcanzaron "
            f"el óptimo de su perfil; con el perfil base, el moveset óptimo cubre {r1['n_tipos_cubiertos']} de 18 tipos. El AG superó "
            "a la búsqueda aleatoria con igual presupuesto, aunque para este tamaño la enumeración exhaustiva es viable y exacta. "
            "Se discuten las limitaciones del modelo, que ignora carga, recarga y fallos condicionales.")


# ---------------------------------------------------------------------------
# PDF con Word
# ---------------------------------------------------------------------------
def export_pdf(docx_path: Path, pdf_path: Path):
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(str(docx_path), AddToRecentFiles=False)
        for i in range(1, doc.TablesOfContents.Count + 1):
            doc.TablesOfContents(i).Update()
        doc.Fields.Update()
        for i in range(1, doc.TablesOfContents.Count + 1):
            doc.TablesOfContents(i).Update()
        doc.Save()
        doc.ExportAsFixedFormat(str(pdf_path), 17)
        doc.Close(False)
    finally:
        word.Quit()


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--repo-url", help="URL publica del repositorio (se guarda en config/report.json)")
    parser.add_argument("--no-pdf", action="store_true")
    args = parser.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    cfg_path = CONFIG_DIR / "report.json"
    if args.repo_url:
        cfg = read_json(cfg_path)
        cfg["repo_url"] = args.repo_url.strip()
        cfg_path.write_text(json.dumps(cfg, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    c = Ctx()
    repo_url = c.report_cfg.get("repo_url")
    builder = build_content(c, repo_url)

    doc = Document(str(DOCS_DIR / "plantilla" / c.report_cfg["plantilla"]))
    fill_template(doc, c, abstract_text(c),
                  "algoritmo genético, DEAP, optimización combinatoria, Pokémon, Charizard, moveset, cobertura de tipos, "
                  "enumeración exhaustiva")
    ensure_outline_levels(doc)
    add_page_number_footer(doc)
    figures = []
    nums = render(doc, builder, figures)
    out_docx = DOCS_DIR / "informe_final.docx"
    doc.save(str(out_docx))

    # Controles del documento generado.
    text = "\n".join(p.text for p in Document(str(out_docx)).paragraphs)
    for t in Document(str(out_docx)).tables:
        for row in t.rows:
            text += "\n" + " ".join(cell.text for cell in row.cells)
    leftovers = [s for s in ("xxx", "xx/xx", "20xx", "Sección xx", "@fig:", "@tab:", "Autor1") if s in text]
    placeholders = text.count("[COMPLETAR")
    checks = {
        "figuras": sum(1 for k in nums if k.startswith("fig:")),
        "tablas": sum(1 for k in nums if k.startswith("tab:")),
        "marcadores_sin_resolver": leftovers,
        "placeholders_completar": placeholders,
        "repo_url": repo_url,
        "figuras_insertadas": [Path(f).relative_to(ROOT).as_posix() for f in figures],
    }
    if leftovers:
        raise SystemExit(f"Quedaron marcadores sin resolver: {leftovers}")
    if not repo_url and placeholders == 0:
        raise SystemExit("Falta el marcador de la URL del repositorio")
    if args.no_pdf:
        print("DOCX generado (sin PDF):", out_docx)
    else:
        out_pdf = DOCS_DIR / "informe_final.pdf"
        export_pdf(out_docx, out_pdf)
        checks["pdf"] = str(out_pdf.relative_to(ROOT))
        print("DOCX y PDF generados:", out_docx.name, out_pdf.name)
    (RESULTS_DIR / "report_checks.json").write_text(json.dumps(checks, indent=2, ensure_ascii=False), encoding="utf-8")
    write_compliance_md(c, repo_url)
    print(f"Figuras: {checks['figuras']}, tablas: {checks['tablas']}, marcadores [COMPLETAR]: {placeholders}")


def write_compliance_md(c: Ctx, repo_url):
    lines = ["# Matriz de cumplimiento", "",
             "Generada por `scripts/build_report.py` a partir de las evidencias (misma fuente que el Anexo A del informe).", "",
             "| Requisito de la consigna / del prompt | Sección del informe | Evidencia | Estado |", "|---|---|---|---|"]
    for a, b, d, e in compliance_rows(c, repo_url):
        lines.append(f"| {a} | {b} | {d} | {e} |")
    lines += ["", "Requisitos exclusivos de redes neuronales (3 entrenamientos, métricas de clasificación/estimación, 20 % de "
              "prueba, datos en CSV/JPG/WAV): **no corresponden** a este proyecto, que implementa un AG.", "",
              "Requisitos adicionales del prompt de trabajo verificados:", "",
              "- Pesos R1–R5 y operadores (ruleta, un punto, mutación 0,20) ejecutados con semillas 42–51: `results/summary_runs.csv`.",
              "- Corridas principales = semilla 42 de R1–R5 (elegida antes de ver resultados): `config/experiments.json`.",
              "- Criterio de calidad definido antes de analizar: `config/quality_criterion.json`.",
              "- Comparación con búsqueda aleatoria, voraz, voraz + búsqueda local y exhaustiva: `results/baselines/`.",
              "- Ejemplo numérico calculado con el mismo código: informe 2.6 y `tests/test_fitness.py`.",
              "- Gráficos generados desde los CSV: `scripts/make_figures.py`.",
              "- Notebook ejecutado: `notebooks/tp2_charizard.ipynb`.",
              "- Guía de defensa: `docs/guia_defensa.md`. Registro de validación: `docs/registro_validacion.md`."]
    (DOCS_DIR / "matriz_cumplimiento.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
