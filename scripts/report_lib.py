"""Utilidades para construir el informe sobre la plantilla oficial con python-docx.

El informe se arma como una lista de bloques (titulos, parrafos, tablas, figuras). Antes de escribir el
documento se numeran figuras y tablas en orden de aparicion y se resuelven las referencias cruzadas
@fig:etiqueta y @tab:etiqueta, de modo que la numeracion es siempre consistente.
"""
from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_PT = 10
TABLE_PT = 8
CAPTION_PT = 9
HEADER_FILL = "E8E7E2"


# ---------------------------------------------------------------------------
# Formato numerico en espanol
# ---------------------------------------------------------------------------
def num(x: float, nd: int = 2) -> str:
    s = f"{x:,.{nd}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def ent(x: int) -> str:
    return f"{int(x):,}".replace(",", ".")


def pct(x: float, nd: int = 1) -> str:
    return num(100 * x, nd) + " %"


# ---------------------------------------------------------------------------
# Bloques
# ---------------------------------------------------------------------------
@dataclass
class Block:
    kind: str
    text: str = ""
    label: str = ""
    data: dict = field(default_factory=dict)


class ReportBuilder:
    def __init__(self):
        self.blocks: list[Block] = []

    def h1(self, text):
        self.blocks.append(Block("h1", text))

    def h2(self, text):
        self.blocks.append(Block("h2", text))

    def h3(self, text):
        self.blocks.append(Block("h3", text))

    def p(self, text, **kw):
        self.blocks.append(Block("p", text, data=kw))

    def bullets(self, items):
        for it in items:
            self.blocks.append(Block("bullet", it))

    def eq(self, text):
        self.blocks.append(Block("eq", text))

    def table(self, label, caption, header, rows, widths_cm, align=None, font_pt=TABLE_PT, bold_first_col=False):
        self.blocks.append(Block("table", caption, label, {"header": header, "rows": rows, "widths": widths_cm,
                                                           "align": align, "font_pt": font_pt, "bold_first_col": bold_first_col}))

    def figure(self, label, path, caption, width_cm=12.0):
        self.blocks.append(Block("figure", caption, label, {"path": str(path), "width": width_cm}))

    def page_break(self):
        self.blocks.append(Block("pagebreak"))

    def toc(self):
        self.blocks.append(Block("toc"))

    def refs(self, items):
        for i, it in enumerate(items, 1):
            self.blocks.append(Block("ref", it, data={"n": i}))

    # Numeracion y referencias cruzadas -----------------------------------------------------
    def numbering(self) -> dict:
        nums, nf, nt = {}, 0, 0
        for b in self.blocks:
            if b.kind == "figure":
                nf += 1
                nums[f"fig:{b.label}"] = f"Fig. {nf}"
            elif b.kind == "table":
                nt += 1
                nums[f"tab:{b.label}"] = f"Tabla {nt}"
        return nums

    def resolve(self, text: str, nums: dict) -> str:
        def repl(m):
            key = m.group(1)
            if key not in nums:
                raise KeyError(f"Referencia cruzada sin destino: {key}")
            return nums[key]
        # Referencias con la forma @fig:etiqueta o @tab:etiqueta (comodas dentro de f-strings).
        return re.sub(r"@((?:fig|tab):[A-Za-z0-9_]+)", repl, text)


# ---------------------------------------------------------------------------
# Escritura sobre la plantilla
# ---------------------------------------------------------------------------
def replace_text_everywhere(doc, old: str, new: str) -> int:
    """Reemplaza texto aunque este partido en varias corridas (incluye cuadros de texto y tablas)."""
    count = 0
    for p in doc.element.body.iter(qn("w:p")):
        ts = list(p.iter(qn("w:t")))
        full = "".join(t.text or "" for t in ts)
        if old in full:
            ts[0].text = full.replace(old, new)
            for t in ts[1:]:
                t.text = ""
            count += 1
    return count


def set_cell_text(cell, text, bold=False, size=11):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Arial"


def ensure_outline_levels(doc):
    """Asigna nivel de esquema a los estilos de titulo de la plantilla para que el indice (TOC \\u) los tome."""
    for name, lvl in (("heading1", 0), ("heading2", 1), ("heading3", 2)):
        st = doc.styles[name]
        ppr = st.element.get_or_add_pPr()
        for old in ppr.findall(qn("w:outlineLvl")):
            ppr.remove(old)
        el = OxmlElement("w:outlineLvl")
        el.set(qn("w:val"), str(lvl))
        ppr.append(el)
        # Los titulos no quedan huerfanos al pie de pagina.
        if ppr.find(qn("w:keepNext")) is None:
            ppr.append(OxmlElement("w:keepNext"))


def add_page_number_footer(doc):
    sec = doc.sections[0]
    footer = sec.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_field(p, "PAGE", "1", size=9)


def _add_field(paragraph, instr: str, placeholder: str, size=None):
    def r_with(child):
        r = OxmlElement("w:r")
        if size:
            rpr = OxmlElement("w:rPr")
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), str(int(size * 2)))
            rpr.append(sz)
            r.append(rpr)
        r.append(child)
        return r

    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = placeholder
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    for child in (b, it, s, t, e):
        paragraph._p.append(r_with(child))


INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|_\{.+?\}|\^\{.+?\}|\[\[HL:.+?\]\]|`[^`]+`)")


def add_rich(paragraph, text: str, size=BODY_PT, color=None, bold_all=False):
    """Mini marcado: **negrita**, *cursiva*, _{subindice}, ^{superindice}, `codigo`, [[HL:texto resaltado]]."""
    for part in INLINE.split(text):
        if not part:
            continue
        kw = {}
        if part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.size = Pt(size * 0.9)
            run.font.name = "Consolas"
            run.bold = bold_all or None
            continue
        if part.startswith("**") and part.endswith("**"):
            part, kw = part[2:-2], {"bold": True}
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, kw = part[1:-1], {"italic": True}
        elif part.startswith("_{"):
            part, kw = part[2:-1], {"sub": True}
        elif part.startswith("^{"):
            part, kw = part[2:-1], {"sup": True}
        elif part.startswith("[[HL:"):
            part, kw = part[5:-2], {"hl": True, "bold": True}
        run = paragraph.add_run(part)
        run.font.size = Pt(size)
        run.bold = True if (kw.get("bold") or bold_all) else None
        run.italic = kw.get("italic")
        if kw.get("sub"):
            run.font.subscript = True
        if kw.get("sup"):
            run.font.superscript = True
        if kw.get("hl"):
            rpr = run._element.get_or_add_rPr()
            hl = OxmlElement("w:highlight")
            hl.set(qn("w:val"), "yellow")
            rpr.append(hl)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return paragraph


def _shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _row_props(row, header=False):
    trpr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trpr.append(cs)
    if header:
        th = OxmlElement("w:tblHeader")
        trpr.append(th)


def _cell_margins(table, top=30, bottom=30, left=60, right=60):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def render(doc, builder: ReportBuilder, figures_check: list):
    nums = builder.numbering()
    for b in builder.blocks:
        if b.kind in ("h1", "h2", "h3"):
            style = {"h1": "heading1", "h2": "heading2", "h3": "heading3"}[b.kind]
            p = doc.add_paragraph(style=style)
            add_rich(p, builder.resolve(b.text, nums), size={"h1": 12, "h2": 10.5, "h3": 10}[b.kind])
        elif b.kind == "p":
            p = doc.add_paragraph(style="p1a")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(4)
            if b.data.get("keep_next"):
                p.paragraph_format.keep_with_next = True
            add_rich(p, builder.resolve(b.text, nums), size=b.data.get("size", BODY_PT))
        elif b.kind == "bullet":
            p = doc.add_paragraph(style="p1a")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, "• " + builder.resolve(b.text, nums), size=BODY_PT)
        elif b.kind == "eq":
            p = doc.add_paragraph(style="equation")
            add_rich(p, builder.resolve(b.text, nums), size=BODY_PT)
            for r in p.runs:
                r.font.name = "Cambria Math"
        elif b.kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif b.kind == "toc":
            p = doc.add_paragraph(style="p1a")
            add_rich(p, "**Índice**", size=12)
            p = doc.add_paragraph(style="p1a")
            _add_field(p, 'TOC \\o "1-2" \\h \\z \\u', "[El índice se actualiza al abrir/exportar el documento]")
        elif b.kind == "ref":
            p = doc.add_paragraph(style="reference")
            add_rich(p, f"{b.data['n']}.\t" + builder.resolve(b.text, nums), size=8.5)
        elif b.kind == "figure":
            path = Path(b.data["path"])
            if not path.exists():
                raise FileNotFoundError(path)
            p = doc.add_paragraph(style="p1a")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
            p.add_run().add_picture(str(path), width=Cm(b.data["width"]))
            cap = doc.add_paragraph(style="figlegend")
            cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_rich(cap, f"**{nums['fig:' + b.label]}.** " + builder.resolve(b.text, nums), size=CAPTION_PT)
            figures_check.append(str(path))
        elif b.kind == "table":
            cap = doc.add_paragraph(style="tablelegend")
            cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            cap.paragraph_format.keep_with_next = True
            add_rich(cap, f"**{nums['tab:' + b.label]}.** " + builder.resolve(b.text, nums), size=CAPTION_PT)
            _render_table(doc, b, builder, nums)
            spacer = doc.add_paragraph(style="p1a")
            spacer.paragraph_format.space_after = Pt(2)
    return nums


def _render_table(doc, b, builder, nums):
    header, rows, widths = b.data["header"], b.data["rows"], b.data["widths"]
    align = b.data["align"] or (["l"] * len(header))
    size = b.data["font_pt"]
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _cell_margins(t)
    for i, row in enumerate([header] + rows):
        tr = t.rows[i]
        _row_props(tr, header=(i == 0))
        for j, val in enumerate(row):
            cell = tr.cells[j]
            cell.width = Cm(widths[j])
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.space_before = Pt(0)
            par.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
                             "r": WD_ALIGN_PARAGRAPH.RIGHT}[align[j] if i > 0 else ("c" if align[j] != "l" else "l")]
            text = builder.resolve(str(val), nums)
            if i == 0:
                add_rich(par, text, size=size, bold_all=True)
                _shade(cell, HEADER_FILL)
            elif j == 0 and b.data["bold_first_col"]:
                add_rich(par, text, size=size, bold_all=True)
            else:
                add_rich(par, text, size=size)
            # Tablas cortas: se mantienen en una sola pagina (cada fila "se mantiene con la siguiente").
            if len(rows) <= 32 and i < len(rows):
                par.paragraph_format.keep_with_next = True
    # Ancho de columnas tambien en la grilla (Word lo respeta mejor).
    grid = t._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[j] * 567)))


def remove_body_from(doc, start_text: str):
    """Elimina los parrafos de ejemplo de la plantilla desde el primero que contiene start_text."""
    body = doc.element.body
    started = False
    for el in list(body.iterchildren()):
        if el.tag == qn("w:sectPr"):
            continue
        txt = "".join(el.itertext())
        if not started and el.tag == qn("w:p") and start_text in txt:
            started = True
        if started:
            body.remove(el)


def duplicate_last_row(table, times: int):
    last = table.rows[-1]._tr
    for _ in range(times):
        last.addnext(copy.deepcopy(last))
